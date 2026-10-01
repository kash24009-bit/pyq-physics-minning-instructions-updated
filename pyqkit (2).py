#!/usr/bin/env python3
"""pyqkit — toolkit for the PYQ Pattern Mining — Physics (042) project.

Standard library + Pillow. PDFs: poppler CLI if present, else PyMuPDF, else pypdfium2 + pypdf.
Excel: openpyxl. Every count, sum, check and derived table in the project comes from here.

Commands: init restore seed inventory pages render merge check-paper similar candidates validate
          derive index year-close export-xlsx plan-mock backtest pack unpack db-export db-import
          status vault-page version            (python3 pyqkit.py <command> -h for arguments)
"""
import argparse, csv, datetime as dt, glob, hashlib, io, json, math, os, random, re, shutil, statistics
import subprocess, sys, tempfile, zipfile
from collections import Counter, defaultdict

KIT_EDITION = "2026-09-30b"
SCHEMA_EDITION = "2026-09-29"
csv.field_size_limit(10**8)

# ------------------------------------------------------------------------------------------------
# Schema — exact column order (P§7, P§10.1). Additive changes only (P§21).
# ------------------------------------------------------------------------------------------------
SCHEMA = {
    "l1/FILES.csv": "file_id sha256 name where fmt kind exam year year_evidence paper_code paper_id sets_covered "
                    "lang pages text_chars visual_pages status batch note",
    "l1/BLUEPRINTS.csv": "paper_id year exam n_questions total_marks sections choice_rule calculators constants_listed "
                         "bp_class reading_time source_page batch",
    "l1/INSTANCES.csv": "instance_id paper_id file_id year exam paper_code q_no side part section marks_part "
                        "q_marks_printed choice_pair_id case_id vi_alt chapter topic_id q_type command_verb has_figure "
                        "wants_figure wants_graph givens unknown options stem source_page read_from rel rel_to "
                        "syllabus_status note extra batch",
    "l1/FIGURES.csv": "figure_id instance_id fig_type printed_or_demanded description values labels_required "
                      "source_page from_image batch",
    "l1/FORMULAE.csv": "formula_id instance_id role formula quantities units source source_page from_image batch",
    "l1/CONSTANTS.csv": "paper_id constant value source_page batch",
    "l1/CASES.csv": "case_id paper_id q_no context chapter n_subparts subpart_types attempt_rule source_page batch",
    "l1/STEP-AWARDS.csv": "award_id instance_id paper_id file_id step_no step_text marks step_kind mandatory "
                          "alt_accepted provenance source_page batch",
    "l1/UNREADABLE.csv": "paper_id page q_no problem what_lost recovered_from batch",
    "l1/ERRATA.csv": "errata_id file row_id field old new reason evidence batch",
    "j/CLASSIFY.csv": "instance_id code family cog_level competency difficulty ncert_ref confidence batch rule_version",
    "j/REGISTRY.csv": "code family chapter topic_id q_type name asks_you_to skeleton traps status alias_of "
                      "issued_batch first_instance note",
    "j/SYLLABUS.csv": "topic_id mark_group group_marks unit chapter chapter_name topic status limitation source",
    "out/MOCKS.csv": "mock_id created_utc mode seed codes sha256 files frozen",
}
SCHEMA = {k: v.split() for k, v in SCHEMA.items()}
ID_COL = {"l1/FILES.csv": "file_id", "l1/BLUEPRINTS.csv": "paper_id", "l1/INSTANCES.csv": "instance_id",
          "l1/FIGURES.csv": "figure_id", "l1/FORMULAE.csv": "formula_id", "l1/CASES.csv": "case_id",
          "l1/STEP-AWARDS.csv": "award_id", "l1/ERRATA.csv": "errata_id", "j/CLASSIFY.csv": "instance_id",
          "j/REGISTRY.csv": "code", "j/SYLLABUS.csv": "topic_id", "out/MOCKS.csv": "mock_id"}
ENUM = {
    "q_type": {"MCQ", "AR", "DEF", "CON", "PRD", "NUM", "DRV", "DGM", "GRF", "CMP", "DEV", "EXP"},
    "exam": {"MAIN", "COMPT", "SQP", "APQ", "T1", "T2", "UNK"},
    "kind": {"QP", "MS", "SQP", "SQP-MS", "APQ", "SYLLABUS", "NCERT", "EXEMPLAR", "COMPILED", "OTHER", "UNKNOWN"},
    "rel": {"origin", "shuffle", "variant", "repeat", ""},
    "syllabus_status": {"IN", "DELETED", "VERIFY-DEL", "X11", "UNK", ""},
    "status": {"NEW", "IN-PROGRESS", "DONE", "QUARANTINE", "DUP-FILE", "UNAVAILABLE"},
    "step_kind": {"formula", "substitution", "result", "unit", "diagram", "label", "reason", "definition", "key", "other"},
    "reg_status": {"DRAFT", "VERIFIED", "ALIAS", "SPLIT"},
    "yn": {"y", "n", ""},
}
EVENT_RELS = {"origin", "repeat", ""}          # shuffle/variant are never counted again
MAIN_LIKE = ("MAIN", "T1", "T2")

# ------------------------------------------------------------------------------------------------
# CSV + state helpers — all writing goes through here (Mistake B2)
# ------------------------------------------------------------------------------------------------
def rpath(vault, rel):
    return os.path.join(vault, rel)

def read_table(vault, rel):
    p = rpath(vault, rel)
    if not os.path.exists(p):
        return list(SCHEMA.get(rel, [])), []
    with open(p, newline="", encoding="utf-8") as f:
        rd = csv.DictReader(f)
        return list(rd.fieldnames or SCHEMA.get(rel, [])), [dict(r) for r in rd]

def write_table(vault, rel, rows, header=None):
    header = header or SCHEMA[rel]
    p = rpath(vault, rel)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    tmp = p + ".tmp"
    with open(tmp, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=header, extrasaction="raise")
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in header})
    os.replace(tmp, p)

def write_csv(path, rows, header=None):
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    header = header or (list(rows[0].keys()) if rows else ["empty"])
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=header, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

def sha256_bytes(b):
    return hashlib.sha256(b).hexdigest()

def sha256_file(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()

def now_utc():
    return dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")

def load_state(vault):
    p = rpath(vault, "STATE.json")
    return json.load(open(p)) if os.path.exists(p) else {}

def save_state(vault, st):
    json.dump(st, open(rpath(vault, "STATE.json"), "w"), indent=2)

def row_counts(vault):
    return {rel: len(read_table(vault, rel)[1]) for rel in SCHEMA}

def apply_errata(vault, rel, rows):
    """Rows with ERRATA applied in order; Layer-1 files themselves are never edited."""
    _, errata = read_table(vault, "l1/ERRATA.csv")
    idc = ID_COL.get(rel)
    if not idc or not errata:
        return rows
    idx = {r[idc]: r for r in rows}
    for e in errata:
        if os.path.basename(e["file"]) == os.path.basename(rel):
            r = idx.get(e["row_id"])
            if r is not None and e["field"] in r:
                r[e["field"]] = e["new"]
    return rows

def table(vault, rel, errata=True):
    _, rows = read_table(vault, rel)
    return apply_errata(vault, rel, rows) if errata else rows

def fnum(x):
    s = str(x if x is not None else "").replace("½", ".5").strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None

def year_int(r):
    y = str(r.get("year", ""))
    return int(y) if y.isdigit() else None

def fam_key(f):
    return (f or "").split(" ")[0]

# ------------------------------------------------------------------------------------------------
# init
# ------------------------------------------------------------------------------------------------
PROGRESS_TEMPLATE = """# PROGRESS — Physics (042) mining
Resume rule: continue from the first unticked line. A ticked item is never redone.

## Done
- Vault created ({date}).

## Next (Claude's work, in order)
- [ ] Batch 0: verify blueprint from SQP 2026-27; verify syllabus; seed; vault round-trip test; mine SQP 2026-27 + scheme.

## Open
- none

## Optional (hers)
- nothing needed
"""

def cmd_init(a):
    v = a.vault
    for sub in ("l1", "j", "out", "l2", "tools", "work", "notes"):
        os.makedirs(rpath(v, sub), exist_ok=True)
    for rel, hdr in SCHEMA.items():
        if not os.path.exists(rpath(v, rel)):
            write_table(v, rel, [], hdr)
    if not os.path.exists(rpath(v, "STATE.json")):
        save_state(v, {"batch": 0, "year_in_progress": "", "years_closed": [], "schema_edition": SCHEMA_EDITION,
                       "kit_edition": KIT_EDITION, "created_utc": now_utc(), "vault_url": "", "rows": row_counts(v)})
    if not os.path.exists(rpath(v, "PROGRESS.md")):
        open(rpath(v, "PROGRESS.md"), "w").write(PROGRESS_TEMPLATE.format(date=dt.date.today().isoformat()))
    if not os.path.exists(rpath(v, "CONVENTIONS.md")):
        open(rpath(v, "CONVENTIONS.md"), "w").write(
            f"# CONVENTIONS\nSchema edition {SCHEMA_EDITION} · kit edition {KIT_EDITION} · created {dt.date.today()}\n"
            "Every convention, blueprint correction, examiner rule and schema change is dated here (P§21).\n")
    if not os.path.exists(rpath(v, "QQUEUE.md")):
        open(rpath(v, "QQUEUE.md"), "w").write("# Q-QUEUE (never blocking)\n")
    me = os.path.abspath(__file__)
    if me != os.path.abspath(rpath(v, "tools/pyqkit.py")):
        shutil.copy2(me, rpath(v, "tools/pyqkit.py"))
    print(f"vault ready at {v} · {len(SCHEMA)} ledgers · schema edition {SCHEMA_EDITION}")

# ------------------------------------------------------------------------------------------------
# Document access — originals (PDF), claude.ai project bundles, images; with fallbacks
# ------------------------------------------------------------------------------------------------
def have_poppler():
    return shutil.which("pdftotext") is not None and not os.environ.get("PYQKIT_NO_POPPLER")

def _fitz():
    try:
        import fitz  # PyMuPDF
        return fitz
    except Exception:
        return None

def _pdfium():
    try:
        import pypdfium2 as pdfium
        return pdfium
    except Exception:
        return None

def is_bundle(p):
    try:
        with zipfile.ZipFile(p) as z:
            return "manifest.json" in z.namelist()
    except (zipfile.BadZipFile, OSError):
        return False

def doc_fmt(p):
    low = p.lower()
    if low.endswith((".jpg", ".jpeg", ".png", ".webp")):
        return "image"
    if is_bundle(p):
        return "bundle"
    try:
        with open(p, "rb") as f:
            head = f.read(5)
    except OSError:
        return "other"
    if head == b"%PDF-":
        return "pdf"
    if zipfile.is_zipfile(p):
        return "zip"
    return "other"

def page_texts(p, fmt=None):
    """Page texts, index 0 = page 1."""
    fmt = fmt or doc_fmt(p)
    if fmt == "bundle":
        with zipfile.ZipFile(p) as z:
            man = json.loads(z.read("manifest.json"))
            names = set(z.namelist())
            out = []
            for pg in man.get("pages", []):
                tp = (pg.get("text") or {}).get("path")
                out.append(z.read(tp).decode("utf-8", "replace") if tp and tp in names else "")
            return out
    if fmt == "pdf":
        if have_poppler():
            r = subprocess.run(["pdftotext", "-layout", p, "-"], capture_output=True)
            pages = r.stdout.decode("utf-8", "replace").split("\f")
            if pages and not pages[-1].strip():
                pages = pages[:-1]
            return pages
        fz = _fitz()
        if fz:
            with fz.open(p) as d:
                return [pg.get_text() for pg in d]
        pdfium = _pdfium()
        if pdfium:
            doc = pdfium.PdfDocument(p)
            out = []
            for i in range(len(doc)):
                tp = doc[i].get_textpage()
                out.append(tp.get_text_range())
            return out
        try:
            from pypdf import PdfReader
            return [pg.extract_text() or "" for pg in PdfReader(p).pages]
        except Exception:
            return [""]
    return [""]

def visual_pages(p, fmt=None):
    fmt = fmt or doc_fmt(p)
    if fmt == "bundle":
        with zipfile.ZipFile(p) as z:
            man = json.loads(z.read("manifest.json"))
            return [pg["page_number"] for pg in man.get("pages", []) if pg.get("has_visual_content")]
    if fmt == "pdf" and shutil.which("pdfimages") and not os.environ.get("PYQKIT_NO_POPPLER"):
        r = subprocess.run(["pdfimages", "-list", p], capture_output=True, text=True)
        pages = set()
        for line in r.stdout.splitlines()[2:]:
            parts = line.split()
            if len(parts) > 4 and parts[0].isdigit():
                try:
                    if int(parts[3]) * int(parts[4]) > 20000:
                        pages.add(int(parts[0]))
                except ValueError:
                    pass
        return sorted(pages)
    if fmt == "pdf":
        fz = _fitz()
        if fz:
            with fz.open(p) as d:
                return [i + 1 for i, pg in enumerate(d) if pg.get_images() or pg.get_drawings()]
    return [1] if fmt == "image" else []

def pdf_meta_year(p):
    if shutil.which("pdfinfo") and not os.environ.get("PYQKIT_NO_POPPLER"):
        r = subprocess.run(["pdfinfo", p], capture_output=True, text=True)
        m = re.search(r"CreationDate:\s+.*?\b((?:19|20)\d\d)\b", r.stdout)
        return m.group(1) if m else ""
    try:
        from pypdf import PdfReader
        md = PdfReader(p).metadata or {}
        m = re.search(r"((?:19|20)\d\d)", str(md.get("/CreationDate", "")))
        return m.group(1) if m else ""
    except Exception:
        return ""

def render_page(p, fmt, pg, dpi):
    """Return a PIL image of page pg (1-based)."""
    from PIL import Image
    if fmt == "bundle":
        with zipfile.ZipFile(p) as z:
            man = json.loads(z.read("manifest.json"))
            im = Image.open(io.BytesIO(z.read(man["pages"][pg - 1]["image"]["path"])))
            im.load()
            return im
    if fmt == "image":
        return Image.open(p)
    if have_poppler():
        with tempfile.TemporaryDirectory() as td:
            base = os.path.join(td, "r")
            subprocess.run(["pdftoppm", "-r", str(dpi), "-f", str(pg), "-l", str(pg), "-gray", "-png", p, base], check=True)
            im = Image.open(sorted(glob.glob(base + "*.png"))[-1])
            im.load()
            return im
    fz = _fitz()
    if fz:
        with fz.open(p) as d:
            pix = d[pg - 1].get_pixmap(dpi=dpi)
            return Image.open(io.BytesIO(pix.tobytes("png")))
    pdfium = _pdfium()
    if pdfium:
        doc = pdfium.PdfDocument(p)
        return doc[pg - 1].render(scale=dpi / 72).to_pil()
    raise RuntimeError("no PDF renderer available (install poppler, PyMuPDF or pypdfium2)")

DEVANAGARI = re.compile(r"[\u0900-\u097F]")
LATIN = re.compile(r"[A-Za-z]")

def page_lang(t):
    d, l = len(DEVANAGARI.findall(t)), len(LATIN.findall(t))
    if d + l < 20:
        return "blank"
    if d > 1.5 * l:
        return "hi"
    if l > 1.5 * d:
        return "en"
    return "bi"

CODE_RE = re.compile(r"(?<![\d/_-])(\d{2})\s*[/_-]\s*([0-9]{1,2}|[A-Z]{1,2})\s*(?:[/_-]\s*([0-9]{1,2}))?((?:\s*,\s*\d)*)(?![\d/_])")
FOOTER_RE = re.compile(r"\b(\d\d)-(\d\d)-(\d\d)N\b")

def guess_code(text):
    zones = re.findall(r"(?:Q\.?\s*P\.?\s*Code|Code|CODE|Paper Code|PAPER CODE|कोड)[^\n]{0,40}", text) + [text]
    for zone in zones:
        for m in CODE_RE.finditer(zone):
            a, b, c, more = m.groups()
            if a in ("55", "56", "65", "30"):
                return f"{a}/{b}" + (f"/{c}" if c else "")
    return ""

def code_sets(text):
    m = CODE_RE.search(text)
    if m and m.group(4) and m.group(1) in ("55", "56", "65", "30"):
        return ",".join([m.group(3) or "?"] + re.findall(r"\d", m.group(4)))
    return ""

def guess_kind(text, name):
    t = text.lower()
    if "marking scheme" in t or "strictly confidential" in t or "value points" in t:
        return "SQP-MS" if "sample" in t else "MS"
    if "sample question paper" in t or "sample paper" in t:
        return "SQP"
    if "additional practice" in t:
        return "APQ"
    if "exemplar" in t:
        return "EXEMPLAR"
    if "curriculum" in t or ("syllabus" in t and "course structure" in t):
        return "SYLLABUS"
    if "general instructions" in t or "q.p. code" in t or "roll no" in t or "series" in t:
        return "QP"
    if "ncert" in name.lower():
        return "NCERT"
    return "UNKNOWN"

def year_evidence(text, name, meta_year, zipname=""):
    cands = []
    for m in re.finditer(r"Examination[s]?\s*[,:\-–]?\s*((?:19|20)\d\d)\b", text, re.I):
        cands.append(("printed", m.group(1)))
    for m in re.finditer(r"(?:Session|Academic Session|Class\s*[-–]?\s*XII)\s*[:\-–]?\s*\(?\s*((?:19|20)\d\d)\s*[-–]\s*(\d\d)", text, re.I):
        cands.append(("printed-session", str(int(m.group(1)) + 1)))
    for m in FOOTER_RE.finditer(text):
        cands.append(("footer", "20" + m.group(1)))
    for src, nm in (("name", name), ("zipname", zipname)):
        m = re.match(r"\s*((?:19|20)\d\d)(?:\s*[-–]\s*((?:19|20)\d\d))?", os.path.basename(nm or ""))
        if m:
            cands.append((src, m.group(1) if not m.group(2) else f"{m.group(1)}-{m.group(2)}"))
    if meta_year:
        cands.append(("pdf-meta", meta_year))
    ev = "; ".join(dict.fromkeys(f"{s}:{y}" for s, y in cands))
    exam_printed = [y for s, y in cands if s == "printed"]
    sess = [y for s, y in cands if s == "printed-session"]
    footer = [y for s, y in cands if s == "footer"]
    named = [y for s, y in cands if s in ("name", "zipname")]
    meta = [y for s, y in cands if s == "pdf-meta"]
    printed = exam_printed or ([max(sess)] if sess else []) or footer
    year, note = "UNK", ""
    if printed:
        year = Counter(printed).most_common(1)[0][0]
        if named and "-" not in named[0] and named[0] != year:
            note = f"name says {named[0]}, page says {year} - page wins"
    elif named and "-" not in named[0]:
        year = named[0]
    elif meta:
        year, note = meta[0], "year from PDF metadata only (weak) - Q-QUEUE"
    return year, ev, note

def guess_exam(text, kind, code):
    t = text.lower()
    if kind in ("SQP", "SQP-MS"):
        return "SQP"
    if kind == "APQ":
        return "APQ"
    if "compartment" in t or "supplementary" in t or re.match(r"\d\d/C", code or ""):
        return "COMPT"
    if re.search(r"term\s*[-–]?\s*(1|i)\b", t):
        return "T1"
    if re.search(r"term\s*[-–]?\s*(2|ii)\b", t):
        return "T2"
    return "MAIN"

def make_paper_id(exam, year, code, sha):
    digits = re.sub(r"[^0-9A-Z]", "", code or "") or ("X" + sha[:6])
    if exam == "SQP":
        y = int(year) if str(year).isdigit() else 0
        return (f"SQP{str(y - 1)[-2:]}{str(y)[-2:]}" if y else "SQPUNK") + "-042"
    tag = {"COMPT": f"{year}C", "T1": f"T1-{year}", "T2": f"T2-{year}"}.get(exam, str(year))
    return f"{tag}-{digits}"

PERMANENT_NAMES = re.compile(r"^(PHY-0\d|PHY-VAULT|PHY-YEAR|PHYSICS-COMMAND)", re.I)

def walk_inputs(paths, intake_dir):
    stack = [(p, "upload", "") for p in paths]
    while stack:
        p, where, zipname = stack.pop(0)
        if os.path.isdir(p):
            for q in sorted(glob.glob(os.path.join(p, "**", "*"), recursive=True)):
                if os.path.isfile(q):
                    stack.append((q, where, zipname))
            continue
        if PERMANENT_NAMES.match(os.path.basename(p)) or p.endswith((".md", ".csv", ".json", ".txt", ".py")):
            continue
        fmt = doc_fmt(p)
        if fmt == "zip":
            dest = os.path.join(intake_dir, re.sub(r"[^\w.-]+", "_", os.path.basename(p)))
            os.makedirs(dest, exist_ok=True)
            with zipfile.ZipFile(p) as z:
                z.extractall(dest)
            for q in sorted(glob.glob(os.path.join(dest, "**", "*"), recursive=True)):
                if os.path.isfile(q):
                    stack.append((q, f"zip:{os.path.basename(p)}", os.path.basename(p)))
            continue
        if fmt in ("pdf", "bundle", "image"):
            yield p, ("project" if p.startswith("/mnt/project") else where), zipname

def cmd_inventory(a):
    v = a.vault
    _, files = read_table(v, "l1/FILES.csv")
    known = {r["sha256"]: r for r in files}
    st = load_state(v)
    batch = str(a.batch if a.batch is not None else st.get("batch", 0))
    added, dupes, report = [], [], []
    for p, where, zipname in walk_inputs(a.paths, rpath(v, "work/_intake")):
        sha = sha256_file(p)
        name = os.path.basename(p)
        if sha in known:
            dupes.append((name, known[sha]["name"]))
            continue
        fmt = doc_fmt(p)
        texts = page_texts(p, fmt)
        head = "\n".join(texts[:3])
        meta = pdf_meta_year(p) if fmt == "pdf" else ""
        kind = guess_kind(head, name)
        code = guess_code(head)
        year, ev, ynote = year_evidence(head, name, meta, zipname)
        exam = guess_exam(head, kind, code)
        langs = Counter(page_lang(t) for t in texts)
        lang = "bi" if langs.get("hi", 0) and langs.get("en", 0) else ("hi" if langs.get("hi", 0) > langs.get("en", 0) else "en")
        notes = [n for n in (ynote,) if n]
        if code and not code.startswith("55") and kind in ("QP", "MS"):
            notes.append(f"not-physics? (code {code})")
        if kind == "UNKNOWN":
            notes.append("kind unknown - check page 1")
        row = {"file_id": "FL-" + sha[:10], "sha256": sha, "name": name, "where": where, "fmt": fmt, "kind": kind,
               "exam": exam, "year": year, "year_evidence": ev, "paper_code": code,
               "paper_id": make_paper_id(exam, year, code, sha) if kind in ("QP", "MS", "SQP", "SQP-MS", "APQ") else "",
               "sets_covered": code_sets(head), "lang": lang, "pages": len(texts),
               "text_chars": sum(len(t.strip()) for t in texts), "visual_pages": " ".join(map(str, visual_pages(p, fmt))),
               "status": "NEW", "batch": batch, "note": "; ".join(notes)}
        known[sha] = row
        added.append(row)
        report.append(f"{row['file_id']}  {kind:7} {exam:5} {year:9} {code or '-':8} {row['paper_id'] or '-':16} "
                      f"{fmt:6} p{row['pages']:<3} {name[:40]}{('  ! ' + row['note']) if row['note'] else ''}")
    write_table(v, "l1/FILES.csv", files + added)
    print(f"inventory: {len(added)} new document(s), {len(dupes)} byte-identical file(s) already known")
    for line in report:
        print("  " + line)
    for n, o in dupes:
        print(f"  = {n} is identical to {o} (recorded once)")

# ------------------------------------------------------------------------------------------------
# pages / render
# ------------------------------------------------------------------------------------------------
def resolve_doc(v, ref):
    if os.path.exists(ref):
        return ref
    files = table(v, "l1/FILES.csv") if v else []
    # kit 2026-09-30 (Q-008): an exact file_id wins; a paper_id shared by a paper and its scheme resolves to the paper
    hits = sorted([r for r in files if ref in (r["file_id"], r["paper_id"], r["name"])],
                  key=lambda r: (r["file_id"] != ref, r["kind"] not in ("QP", "SQP", "APQ")))
    for r in hits:
        if True:
            for base in ("/mnt/user-data/uploads", "/mnt/project", rpath(v, "work/_intake")):
                hits = glob.glob(os.path.join(base, "**", r["name"]), recursive=True)
                if hits:
                    return hits[0]
    sys.exit(f"cannot find document {ref}")

def cmd_pages(a):
    p = resolve_doc(a.vault, a.doc)
    fmt = doc_fmt(p)
    texts = page_texts(p, fmt)
    if a.text:
        print(texts[a.text - 1] if 0 < a.text <= len(texts) else "(no such page)")
        return
    vis = set(visual_pages(p, fmt))
    print(f"{os.path.basename(p)} · {fmt} · {len(texts)} pages · English pages to view: "
          + " ".join(str(i + 1) for i, t in enumerate(texts) if page_lang(t) in ("en", "bi", "blank")))
    for i, t in enumerate(texts, 1):
        print(f"  p{i:<3} {page_lang(t):5} chars={len(t.strip()):<5} visual={'y' if i in vis else 'n'}  "
              + re.sub(r"\s+", " ", t.strip())[:70])

def parse_pages(spec, n):
    out = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if part == "all":
            return list(range(1, n + 1))
        if "-" in part:
            x, y = part.split("-")
            out += list(range(int(x), int(y) + 1))
        else:
            out.append(int(part))
    return [q for q in out if 1 <= q <= n]

def crop_margins(im):
    from PIL import ImageOps
    g = ImageOps.invert(im.convert("L"))
    box = g.point(lambda x: 255 if x > 40 else 0).getbbox()
    if box:
        pad = 12
        return im.crop((max(0, box[0] - pad), max(0, box[1] - pad), min(im.width, box[2] + pad), min(im.height, box[3] + pad)))
    return im

def cmd_render(a):
    from PIL import Image
    p = resolve_doc(a.vault, a.doc)
    fmt = doc_fmt(p)
    os.makedirs(a.out, exist_ok=True)
    n = len(page_texts(p, fmt)) if fmt != "image" else 1
    imgs = []
    for pg in parse_pages(a.pages, n):
        im = render_page(p, fmt, pg, a.dpi)
        if a.crop:
            im = crop_margins(im)
        if a.region:
            x0, y0, x1, y1 = [float(t) for t in a.region.split(",")]
            im = im.crop((int(x0 * im.width), int(y0 * im.height), int(x1 * im.width), int(y1 * im.height)))
        imgs.append((pg, im.convert("L")))
    outs = []
    if a.pair:
        for i in range(0, len(imgs), 2):
            grp = imgs[i:i + 2]
            w = sum(im.width for _, im in grp) + 10 * (len(grp) - 1)
            canvas = Image.new("L", (w, max(im.height for _, im in grp)), 255)
            x = 0
            for _, im in grp:
                canvas.paste(im, (x, 0))
                x += im.width + 10
            fn = os.path.join(a.out, f"p{grp[0][0]}" + (f"-{grp[-1][0]}" if len(grp) > 1 else "") + ".png")
            canvas.save(fn)
            outs.append(fn)
    else:
        for pg, im in imgs:
            fn = os.path.join(a.out, f"p{pg}.png")
            im.save(fn)
            outs.append(fn)
    print("\n".join(outs))

# ------------------------------------------------------------------------------------------------
# merge — append a part file into a ledger (append-only; never overwrite)
# ------------------------------------------------------------------------------------------------
def resolve_table(name):
    if "/" in name:
        return name
    for k in SCHEMA:
        if k.endswith("/" + name + ".csv") or k.endswith("/" + name):
            return k
    sys.exit(f"unknown table {name}")

def cmd_merge(a):
    v, rel = a.vault, resolve_table(a.table)
    hdr = SCHEMA[rel]
    with open(a.part, newline="", encoding="utf-8") as f:
        rd = csv.DictReader(f)
        missing = [c for c in hdr if c not in (rd.fieldnames or [])]
        extra = [c for c in (rd.fieldnames or []) if c not in hdr]
        if missing or extra:
            sys.exit(f"REFUSED: header mismatch for {rel}. missing={missing} extra={extra}")
        new = [dict(r) for r in rd]
    _, rows = read_table(v, rel)
    idc = ID_COL.get(rel)
    if idc and rel not in ("l1/BLUEPRINTS.csv",):
        have = {r[idc] for r in rows}
        clash = [r[idc] for r in new if r[idc] in have]
        dup = [k for k, c in Counter(r[idc] for r in new).items() if c > 1]
        if clash or dup:
            sys.exit(f"REFUSED: {len(clash)} id(s) already in {rel} (use ERRATA) and {len(dup)} duplicated "
                     f"inside the part file: {(clash + dup)[:5]}")
    write_table(v, rel, rows + new)
    print(f"merged {len(new)} row(s) into {rel} · now {len(rows) + len(new)}")

# ------------------------------------------------------------------------------------------------
# seed — syllabus + DRAFT archetype catalogue (never overwrites)
# ------------------------------------------------------------------------------------------------
def cmd_seed(a):
    v = a.vault
    added = {}
    for rel, src in (("j/SYLLABUS.csv", a.syllabus), ("j/REGISTRY.csv", a.archetypes)):
        if not src:
            continue
        if not os.path.exists(src):
            print(f"  skip {rel}: {src} not found")
            continue
        _, rows = read_table(v, rel)
        have = {r[ID_COL[rel]] for r in rows}
        with open(src, newline="", encoding="utf-8") as f:
            new = []
            for r in csv.DictReader(f):
                r = {k: r.get(k, "") for k in SCHEMA[rel]}
                if r[ID_COL[rel]] and r[ID_COL[rel]] not in have:
                    new.append(r)
                    have.add(r[ID_COL[rel]])
        write_table(v, rel, rows + new)
        added[rel] = len(new)
    print("seed: " + ", ".join(f"{k} +{n}" for k, n in added.items()))

# ------------------------------------------------------------------------------------------------
# paper checks (V04, V05, V06, V12)
# ------------------------------------------------------------------------------------------------
def parse_sections(spec):
    out = {}
    for part in (spec or "").split(";"):
        part = part.strip()
        if not part:
            continue
        sec, rng, mk = part.split(":")
        x, y = rng.split("-") if "-" in rng else (rng, rng)
        for q in range(int(x), int(y) + 1):
            out[q] = (sec, float(mk))
    return out

def check_paper(v, pid, inst=None, bps=None):
    inst = inst if inst is not None else table(v, "l1/INSTANCES.csv")
    bps = bps if bps is not None else {r["paper_id"]: r for r in table(v, "l1/BLUEPRINTS.csv")}
    rows = [r for r in inst if r["paper_id"] == pid and r.get("vi_alt", "") != "y"]
    probs, info = [], []
    bp = bps.get(pid)
    if not rows:
        return ["no instance rows"], info
    if not bp:
        return ["no BLUEPRINTS row"], info
    n = int(fnum(bp["n_questions"]) or 0)
    qnos = sorted({int(r["q_no"]) for r in rows if str(r["q_no"]).isdigit()})
    missing = [q for q in range(1, n + 1) if q not in qnos]
    if missing:
        probs.append(f"questions missing: {missing}")
    beyond = [q for q in qnos if q > n]
    if beyond:
        probs.append(f"question numbers beyond {n}: {beyond}")
    secmap = parse_sections(bp.get("sections", ""))
    qmarks, sectot = {}, defaultdict(float)
    for q in qnos:
        qrows = [r for r in rows if str(r["q_no"]) == str(q)]
        sides = sorted({r.get("side", "") for r in qrows})
        # kit 2026-09-30: when a question has side-blank rows AND sides A/B (a choice inside one sub-part,
        # 2026 case studies), the side-blank rows are common to both alternatives and count with each side.
        common = [r for r in qrows if r.get("side", "") == ""] if ("" in sides and len(sides) > 1) else []
        side_tot = {}
        for s in sides:
            if s == "" and common:
                continue
            srows = [r for r in qrows if r.get("side", "") == s] + common
            printed = {fnum(r["q_marks_printed"]) for r in srows if fnum(r["q_marks_printed"]) is not None}
            parts = [fnum(r["marks_part"]) for r in srows]
            known = [x for x in parts if x is not None]
            if len(printed) > 1:
                probs.append(f"Q{q}{s}: rows disagree on q_marks_printed {sorted(printed)}")
            tot = next(iter(printed)) if printed else (sum(known) if known and len(known) == len(parts) else None)
            if tot is None and q in secmap:
                tot = secmap[q][1]
                info.append(f"Q{q}{s}: marks from blueprint (none printed)")
            if known and tot is not None and sum(known) - tot > 1e-9:
                probs.append(f"Q{q}{s}: part marks {sum(known):g} exceed printed total {tot:g}")
            elif known and len(known) == len(parts) and tot is not None and abs(sum(known) - tot) > 1e-9:
                probs.append(f"Q{q}{s}: part marks sum {sum(known):g} differs from printed {tot:g}")
            side_tot[s] = tot
        vals = {x for x in side_tot.values() if x is not None}
        if len(vals) > 1:
            probs.append(f"Q{q}: OR sides carry different marks {side_tot}")
        cps = {r.get("choice_pair_id", "") for r in qrows if r.get("choice_pair_id", "")}
        if len([s for s in sides if s]) >= 2 and len(cps) != 1:
            probs.append(f"Q{q}: sides A/B present but choice_pair_id is {sorted(cps) or 'missing'}")
        if cps and not all(c.startswith("OR-") for c in cps):
            probs.append(f"Q{q}: choice_pair_id without OR- prefix")
        qm = next(iter(vals)) if vals else None
        qmarks[q] = qm
        if q in secmap and qm is not None:
            sec, expect = secmap[q]
            sectot[sec] += qm
            if abs(qm - expect) > 1e-9:
                probs.append(f"Q{q}: {qm:g} marks but blueprint says {expect:g}")
    total = sum(x for x in qmarks.values() if x is not None)
    tm = fnum(bp["total_marks"])
    if tm is not None and abs(total - tm) > 1e-9:
        probs.append(f"paper total {total:g} differs from blueprint {tm:g} (OR counted once, VI excluded)")
    if secmap:
        exp = defaultdict(float)
        for q, (s, m) in secmap.items():
            exp[s] += m
        for s in exp:
            if abs(sectot.get(s, 0) - exp[s]) > 1e-9:
                probs.append(f"section {s}: {sectot.get(s, 0):g} vs {exp[s]:g}")
    info.append(f"{len(qnos)}/{n} questions · total {total:g}/{(tm or 0):g}")
    return probs, info

def cmd_check_paper(a):
    probs, info = check_paper(a.vault, a.paper_id)
    print(("PASS" if not probs else "QUARANTINE") + f" · {a.paper_id} · " + "; ".join(info))
    for p in probs:
        print("  x " + p)
    sys.exit(0 if not probs else 2)

# ------------------------------------------------------------------------------------------------
# similar — shuffle / variant / repeat suggestions (Claude confirms on images)
# ------------------------------------------------------------------------------------------------
NUM_RE = re.compile(r"\d+(?:\.\d+)?(?:e-?\d+)?")
STOP = set("the a an of to in is and or for with on by at as be are was it its this that from which what how if when "
           "then than into find calculate given show".split())

def norm_tokens(text):
    s = NUM_RE.sub(" ", (text or "").lower())
    return {t for t in re.findall(r"[a-z]+", s) if t not in STOP and len(t) > 1}

def numbers(text):
    return sorted(NUM_RE.findall((text or "").lower()))

def series_of(code):
    parts = (code or "").split("/")
    return "/".join(parts[:2]) if len(parts) >= 2 else (code or "")

def row_series(r):
    """kit 2026-09-30b: a row may name its series in `extra` as `series=55/3`. Letter-code visually-impaired papers
    (paper_code 55/B, 55/B/7) belong to the series printed on their own page 1 (CONVENTIONS B01: 55(B) is set 5 of
    55/3, 55(B)/7 is set 5 of 55/7), which their paper code cannot show. Without the key, the paper code decides."""
    m = re.search(r"(?:^|;)\s*series\s*=\s*([^;]+)", (r.get("extra") or ""))
    return m.group(1).strip() if m else series_of(r.get("paper_code", ""))

def cmd_similar(a):
    v = a.vault
    inst = [r for r in table(v, "l1/INSTANCES.csv") if r.get("vi_alt") != "y" and len(norm_tokens(r["stem"])) >= 4]
    if a.year and not a.cross:
        inst = [r for r in inst if r["year"] == str(a.year)]
    toks = {r["instance_id"]: norm_tokens(r["stem"]) for r in inst}
    df = Counter(t for s in toks.values() for t in s)
    index = defaultdict(set)
    for iid, s in toks.items():
        for t in sorted(s, key=lambda t: df[t])[:6]:
            index[t].add(iid)
    by_id = {r["instance_id"]: r for r in inst}
    out, seen = [], set()
    for iid, s in toks.items():
        r = by_id[iid]
        cands = Counter()
        for t in sorted(s, key=lambda t: df[t])[:6]:
            for j in index[t]:
                if j != iid:
                    cands[j] += 1
        for j, c in cands.items():
            if c < 2 or (j, iid) in seen:
                continue
            q = by_id[j]
            if q["paper_id"] == r["paper_id"]:
                continue
            seen.add((iid, j))
            union = len(s | toks[j])
            score = len(s & toks[j]) / union if union else 0
            same_series = (r["year"] == q["year"] and row_series(r) == row_series(q)
                           and fnum(r["q_marks_printed"]) == fnum(q["q_marks_printed"]))
            same_nums = numbers(r["stem"]) == numbers(q["stem"])
            if same_series and score >= 0.85 and same_nums:
                sug = "shuffle"
            elif same_series and score >= 0.7:
                sug = "variant"
            elif score >= 0.8:
                sug = "repeat"
            else:
                continue
            a_, b_ = sorted([r, q], key=lambda x: (x["batch"], x["instance_id"]))
            out.append({"instance_id": b_["instance_id"], "rel_to": a_["instance_id"], "suggest": sug,
                        "score": f"{score:.2f}", "numbers_equal": "y" if same_nums else "n",
                        "_rank": (score, a_.get("part", "") == b_.get("part", ""), a_["q_no"] == b_["q_no"])})
    best = {}
    for x in out:
        if x["instance_id"] not in best or x["_rank"] > best[x["instance_id"]]["_rank"]:
            best[x["instance_id"]] = x
    rows = [{k: w for k, w in x.items() if k != "_rank"} for x in best.values()]
    write_csv(rpath(v, "l2/SIMILAR-SUGGESTIONS.csv"), sorted(rows, key=lambda x: (x["suggest"], x["instance_id"])),
              ["instance_id", "rel_to", "suggest", "score", "numbers_equal"])
    c = Counter(x["suggest"] for x in rows)
    print(f"similar: {len(rows)} suggestion(s) - shuffle {c['shuffle']}, variant {c['variant']}, repeat {c['repeat']}"
          f" -> l2/SIMILAR-SUGGESTIONS.csv (confirm each on the page images)")

# ------------------------------------------------------------------------------------------------
# candidates — consistent classification across chats
# ------------------------------------------------------------------------------------------------
def cmd_candidates(a):
    v = a.vault
    reg = [r for r in table(v, "j/REGISTRY.csv") if r["status"] in ("DRAFT", "VERIFIED")]
    inst = table(v, "l1/INSTANCES.csv")
    cls = {r["instance_id"]: r["code"] for r in table(v, "j/CLASSIFY.csv")}
    if a.instance:
        me = next((r for r in inst if r["instance_id"] == a.instance), None)
        if not me:
            sys.exit(f"no instance {a.instance}")
        text, chapter, qtype = me["stem"], a.chapter or me["chapter"], me["q_type"]
    else:
        text, chapter, qtype = a.text or "", a.chapter or "", a.q_type or ""
    q = norm_tokens(text)
    stems = defaultdict(set)
    for r in inst:
        c = cls.get(r["instance_id"])
        if c:
            stems[c] |= norm_tokens(r["stem"])
    scored = []
    for r in reg:
        if chapter and r["chapter"] != chapter and not a.any_chapter:
            continue
        toks = norm_tokens(r["name"] + " " + r["asks_you_to"] + " " + r["skeleton"]) | stems.get(r["code"], set())
        union = len(q | toks)
        s = len(q & toks) / union if union else 0
        s += 0.10 if qtype and r["q_type"] == qtype else 0
        s += 0.05 if r["status"] == "VERIFIED" else 0
        scored.append((s, r))
    scored.sort(key=lambda x: -x[0])
    print(f"candidates for {a.instance or 'text'} (chapter {chapter or 'any'}, type {qtype or 'any'}):")
    for s, r in scored[:a.k]:
        print(f"  {s:.2f}  {r['code']}  [{r['status']}] {r['q_type']}  {r['name']}")
    if not scored:
        print("  none - issue a new code (P§10.3)")

# ------------------------------------------------------------------------------------------------
# validate (P§12)
# ------------------------------------------------------------------------------------------------
ID_OK = re.compile(r"^[A-Za-z0-9]+(?:-[A-Za-z0-9]+)*$")

def validate(v):
    res, quarantine = [], set()
    def add(code, name, fails, detail=""):
        res.append((code, name, "PASS" if not fails else "FAIL",
                    detail if not fails else "; ".join(map(str, fails[:6])) + (f" (+{len(fails) - 6} more)" if len(fails) > 6 else "")))
    bad = [rel for rel, hdr in SCHEMA.items() if os.path.exists(rpath(v, rel)) and read_table(v, rel)[0] != hdr]
    add("V01", "ledger headers match schema", bad, f"{len(SCHEMA)} files")
    inst = table(v, "l1/INSTANCES.csv")
    files = table(v, "l1/FILES.csv")
    bps = {r["paper_id"]: r for r in table(v, "l1/BLUEPRINTS.csv")}
    ids = [r["instance_id"] for r in inst]
    dup = [k for k, c in Counter(ids).items() if c > 1]
    malformed = [i for i in ids if not ID_OK.match(i) or "." in i]
    add("V02", "instance ids unique, well-formed, no dots", dup + malformed, f"{len(ids)} ids")
    papers_known = {r["paper_id"] for r in files if r["paper_id"]}
    v03 = []
    for r in inst:
        if r["paper_id"] not in papers_known:
            v03.append(f"{r['instance_id']}: paper not in FILES")
        if not r["year"]:
            v03.append(f"{r['instance_id']}: no year")
        if not str(r["source_page"]).isdigit():
            v03.append(f"{r['instance_id']}: source_page '{r['source_page']}'")
    add("V03", "rows have paper in FILES, year, source page", v03, f"{len(inst)} rows")
    pids = sorted({r["paper_id"] for r in inst})
    v04, v05 = [], []
    for pid in pids:
        probs, _ = check_paper(v, pid, inst, bps)
        for p in probs:
            (v05 if p.startswith("questions missing") or "beyond" in p else v04).append(f"{pid}: {p}")
            quarantine.add(pid)
    add("V04", "marks = blueprint (OR once, VI excluded), per section", v04, f"{len(pids)} papers")
    add("V05", "question numbers continuous", v05, f"{len(pids)} papers")
    cp = defaultdict(set)
    for r in inst:
        if r.get("choice_pair_id"):
            cp[r["choice_pair_id"]].add(r.get("side", ""))
    v06 = [f"{k}: sides {sorted(s)}" for k, s in cp.items() if (s - {""}) != {"A", "B"} or not k.startswith("OR-")]
    add("V06", "OR pairs have sides A and B and OR- prefix", v06, f"{len(cp)} pairs")
    byid = {r["instance_id"]: r for r in inst}
    v07 = []
    for r in inst:
        rel = r.get("rel", "")
        if rel not in ENUM["rel"]:
            v07.append(f"{r['instance_id']}: rel '{rel}'")
        if rel in ("shuffle", "variant", "repeat"):
            t = byid.get(r.get("rel_to", ""))
            if not t:
                v07.append(f"{r['instance_id']}: rel_to missing")
            elif rel in ("shuffle", "variant") and (t["year"] != r["year"] or row_series(t) != row_series(r)):  # kit 2026-09-30b
                v07.append(f"{r['instance_id']}: {rel} across year/series")
    add("V07", "relations valid (shuffle/variant within year+series)", v07)
    v08 = [f"{os.path.basename(rel)} {r.get(ID_COL[rel])}" for rel in ("l1/FIGURES.csv", "l1/FORMULAE.csv")
           for r in table(v, rel) if r.get("from_image") != "y"]
    add("V08", "figures and formulas read from images", v08)
    steps = table(v, "l1/STEP-AWARDS.csv")
    v09, prov, summ = [], defaultdict(set), defaultdict(float)
    qm = {r["instance_id"]: fnum(r["marks_part"]) if fnum(r["marks_part"]) is not None else fnum(r["q_marks_printed"]) for r in inst}
    for s in steps:
        p = s.get("provenance", "")
        kind = "TRANSFERRED" if p.startswith("TRANSFERRED:") else p
        if kind not in ("PRINTED", "TRANSFERRED", "INFERRED"):
            v09.append(f"{s['award_id']}: provenance '{p}'")
        prov[s["instance_id"]].add(kind)
        if s["instance_id"] not in byid:
            v09.append(f"{s['award_id']}: unknown instance")
        if kind == "PRINTED" and fnum(s["marks"]) is not None:
            summ[s["instance_id"]] += fnum(s["marks"])
    v09 += [f"{i}: mixed provenance {sorted(k)}" for i, k in prov.items() if len(k) > 1]
    v09 += [f"{i}: printed steps {t:g} > marks {qm.get(i)}" for i, t in summ.items() if qm.get(i) is not None and t - qm[i] > 1e-9]
    add("V09", "step awards: provenance, no mixing, <= marks", v09, f"{len(steps)} steps")
    reg = table(v, "j/REGISTRY.csv")
    codes = Counter(r["code"] for r in reg)
    regd = {r["code"]: r for r in reg}
    classify = table(v, "j/CLASSIFY.csv")
    v10 = [f"duplicate code {c}" for c, n in codes.items() if n > 1]
    v10 += [f"{r['instance_id']}: code {r['code']} not in REGISTRY" for r in classify if r["code"] and r["code"] not in codes]
    v10 += [f"{r['code']}: alias_of {r['alias_of']} missing" for r in reg if r["status"] == "ALIAS" and r["alias_of"] not in codes]
    add("V10", "classification codes exist in REGISTRY", v10, f"{len(codes)} codes")
    v11 = []
    for r in inst:
        if r["q_type"] and r["q_type"] not in ENUM["q_type"]:
            v11.append(f"{r['instance_id']}: q_type {r['q_type']}")
        if r["exam"] not in ENUM["exam"]:
            v11.append(f"{r['instance_id']}: exam {r['exam']}")
        if r["syllabus_status"] not in ENUM["syllabus_status"]:
            v11.append(f"{r['instance_id']}: syllabus_status {r['syllabus_status']}")
        for c in ("vi_alt", "has_figure", "wants_figure", "wants_graph"):
            if r.get(c, "") not in ENUM["yn"]:
                v11.append(f"{r['instance_id']}: {c}={r[c]}")
    for r in files:
        if r["status"] not in ENUM["status"]:
            v11.append(f"{r['file_id']}: status {r['status']}")
        if r["kind"] not in ENUM["kind"]:
            v11.append(f"{r['file_id']}: kind {r['kind']}")
    v11 += [f"{s['award_id']}: step_kind {s['step_kind']}" for s in steps if s["step_kind"] and s["step_kind"] not in ENUM["step_kind"]]
    v11 += [f"{r['code']}: status {r['status']}" for r in reg if r["status"] not in ENUM["reg_status"]]
    add("V11", "enum values valid", v11)
    add("V12", "part shares <= printed group total", [x for x in v04 if "exceed" in x])
    shas = Counter(r["sha256"] for r in files if r["status"] != "DUP-FILE")
    add("V13", "no duplicate files", [f"sha {k[:10]} x{n}" for k, n in shas.items() if n > 1], f"{len(files)} files")
    v14 = sorted({r["code"] for r in classify if r["code"] in regd and regd[r["code"]]["status"] != "VERIFIED"})
    add("V14", "every code in use is VERIFIED", [f"{c} is {regd[c]['status']}" for c in v14])
    syl = {r["topic_id"] for r in table(v, "j/SYLLABUS.csv")}
    v15 = [f"{r['instance_id']}: topic {r['topic_id']}" for r in inst if r["topic_id"] not in ("", "UNK") and syl and r["topic_id"] not in syl]
    add("V15", "topic ids exist in SYLLABUS", v15)
    add("V16", "report numbers computed by code", [], "by construction")
    q_files = {r["paper_id"] for r in files if r["status"] == "QUARANTINE"}
    return res, sorted(quarantine | q_files)

def cmd_validate(a):
    res, quarantine = validate(a.vault)
    lines = ["# VALIDATION", f"Computed by pyqkit {KIT_EDITION} at {now_utc()}", "",
             "| # | Check | Result | Detail |", "|---|---|---|---|"]
    lines += [f"| {c} | {n} | {r} | {d.replace('|', '/')} |" for c, n, r, d in res]
    passed = sum(1 for x in res if x[2] == "PASS")
    lines += ["", f"**{passed}/{len(res)} checks pass.** Quarantined papers: {', '.join(quarantine) if quarantine else 'none'}."]
    os.makedirs(rpath(a.vault, "l2"), exist_ok=True)
    open(rpath(a.vault, "l2/VALIDATION.md"), "w").write("\n".join(lines) + "\n")
    print("\n".join(lines[3:]))
    sys.exit(0 if passed == len(res) else 1)

# ------------------------------------------------------------------------------------------------
# statistics (P§15.2) — PRESENCE scope, main exams, events only
# ------------------------------------------------------------------------------------------------
def canon_map(v):
    reg = table(v, "j/REGISTRY.csv")
    alias = {r["code"]: r["alias_of"] for r in reg if r["status"] == "ALIAS" and r["alias_of"]}
    def canon(c):
        seen = set()
        while c in alias and c not in seen:
            seen.add(c)
            c = alias[c]
        return c
    return canon, {r["code"]: r for r in reg}

def events(inst, exams=MAIN_LIKE):
    return [r for r in inst if r.get("rel", "") in EVENT_RELS and r.get("vi_alt") != "y" and r["exam"] in exams and year_int(r)]

def archetype_stats(v, until=None, target=None, quarantine=()):
    q = set(quarantine)
    inst = [r for r in table(v, "l1/INSTANCES.csv") if r["paper_id"] not in q]
    canon, reg = canon_map(v)
    cls = {r["instance_id"]: canon(r["code"]) for r in table(v, "j/CLASSIFY.csv") if r["code"]}
    ev = [r for r in events(inst) if until is None or year_int(r) <= until]
    years_all = sorted({year_int(r) for r in ev})
    target = target or ((until + 1) if until else (max(years_all) + 1 if years_all else 2027))
    papers_by_year = defaultdict(set)
    term_ch = defaultdict(set)
    for r in ev:
        papers_by_year[year_int(r)].add(r["paper_id"])
        if r["exam"] in ("T1", "T2"):
            term_ch[year_int(r)].add(r["chapter"])
    sqp = defaultdict(set)
    for r in inst:
        if r["exam"] == "SQP" and r["instance_id"] in cls and year_int(r):
            sqp[year_int(r)].add(cls[r["instance_id"]])
    per = defaultdict(lambda: {"years": set(), "events": set(), "marks": Counter(), "sections": Counter(), "papers": defaultdict(set)})
    for r in ev:
        c = cls.get(r["instance_id"])
        if not c:
            continue
        y = year_int(r)
        d = per[c]
        d["years"].add(y)
        d["events"].add((r["paper_id"], r["q_no"], r.get("side", "")))
        m = fnum(r["q_marks_printed"])
        if m is not None:
            d["marks"][f"{m:g}"] += 1
        d["sections"][r["section"]] += 1
        d["papers"][y].add(r["paper_id"])
    rows = []
    for code, d in per.items():
        rg = reg.get(code, {})
        ch = rg.get("chapter", "")
        elig = [y for y in years_all if not (y in term_ch and ch and ch not in term_ch[y])]
        ys = sorted(d["years"])
        gaps = [b - a for a, b in zip(ys, ys[1:])]
        mg = statistics.median(gaps) if gaps else None
        w = lambda y: 0.5 ** ((target - y) / 3.0)
        rec = (sum(w(y) for y in ys) / sum(w(y) for y in elig)) if elig else 0
        due = min(2.0, (target - ys[-1]) / mg) if mg else (1.0 if ys else 0)
        spread = statistics.mean(len(d["papers"][y]) / max(1, len(papers_by_year[y])) for y in ys) if ys else 0
        sq = 1 if code in sqp.get(target, set()) else 0
        score = 0.5 * rec + 0.2 * (due / 2) + 0.2 * sq + 0.1 * spread
        rows.append({"code": code, "name": rg.get("name", ""), "chapter": ch, "q_type": rg.get("q_type", ""),
                     "family": rg.get("family", ""), "status": rg.get("status", ""), "years": " ".join(map(str, ys)),
                     "n_years": len(ys), "eligible_years": len(elig), "p_year": f"{len(ys) / len(elig):.2f}" if elig else "",
                     "events": len(d["events"]), "marks_seen": " ".join(f"{k}x{n}" for k, n in sorted(d["marks"].items())),
                     "sections_seen": " ".join(f"{k}x{n}" for k, n in sorted(d["sections"].items())),
                     "first_seen": ys[0] if ys else "", "last_seen": ys[-1] if ys else "", "median_gap": f"{mg:g}" if mg else "",
                     "recency": f"{rec:.3f}", "due": f"{due:.2f}", "sqp_target": sq, "spread": f"{spread:.2f}",
                     "score_uncalibrated": f"{score:.3f}"})
    rows.sort(key=lambda r: -float(r["score_uncalibrated"]))
    return rows, target

# ------------------------------------------------------------------------------------------------
# derive — rebuild all of l2 (never patch)
# ------------------------------------------------------------------------------------------------
REGISTERS = [
    ("R01-derivations", lambda r: r["q_type"] == "DRV"),
    ("R02-numericals", lambda r: r["q_type"] == "NUM"),
    ("R03-reasons", lambda r: r["q_type"] == "CON"),
    ("R04-what-happens-if", lambda r: r["q_type"] == "PRD"),
    ("R05-definitions-laws-units", lambda r: r["q_type"] == "DEF"),
    ("R06-diagrams", lambda r: r["q_type"] == "DGM" or r.get("wants_figure") == "y"),
    ("R07-graphs", lambda r: r["q_type"] == "GRF" or r.get("wants_graph") == "y"),
    ("R08-devices", lambda r: r["q_type"] == "DEV"),
    ("R09-mcq", lambda r: r["q_type"] == "MCQ"),
    ("R10-assertion-reason", lambda r: r["q_type"] == "AR"),
]
BASE_COLS = ["code", "instance_id", "year", "exam", "paper_code", "q_no", "side", "part", "q_marks_printed",
             "chapter", "q_type", "command_verb", "stem", "rel"]

def md_table(rows, cols, limit=None):
    if not rows:
        return "_none yet_\n"
    out = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for r in rows[:limit] if limit else rows:
        out.append("| " + " | ".join(str(r.get(c, "")).replace("|", "/").replace("\n", " ")[:160] for c in cols) + " |")
    return "\n".join(out) + "\n"

def chapter_names(v):
    return {r["chapter"]: r["chapter_name"] for r in table(v, "j/SYLLABUS.csv")}

def cmd_derive(a):
    v = a.vault
    l2 = rpath(v, "l2")
    _, quarantine = validate(v)
    if a.until:
        stats, target = archetype_stats(v, until=a.until, quarantine=quarantine)
        out = os.path.join(l2, "backtest", f"ARCHETYPE-STATS-until-{a.until}.csv")
        write_csv(out, stats)
        print(f"derive --until {a.until}: {len(stats)} archetypes scored for {target} -> {out}")
        return
    for f in glob.glob(os.path.join(l2, "**", "*"), recursive=True):
        if os.path.isfile(f) and not f.endswith(("VALIDATION.md", "SIMILAR-SUGGESTIONS.csv", "BACKTEST.md")) and "/backtest/" not in f:
            os.remove(f)
    q = set(quarantine)
    inst = table(v, "l1/INSTANCES.csv")
    live = [r for r in inst if r["paper_id"] not in q]
    files = table(v, "l1/FILES.csv")
    canon, reg = canon_map(v)
    cls = {r["instance_id"]: r for r in table(v, "j/CLASSIFY.csv")}
    code_of = lambda r: canon(cls.get(r["instance_id"], {}).get("code", "")) if cls.get(r["instance_id"], {}).get("code") else ""
    syl = table(v, "j/SYLLABUS.csv")
    chn = chapter_names(v)
    steps = table(v, "l1/STEP-AWARDS.csv")
    # counts
    by_exam = Counter((r["year"], r["exam"]) for r in {r["paper_id"]: r for r in live}.values())
    ev = events(live)
    lines = ["# COUNTS (computed)", f"pyqkit {KIT_EDITION} · {now_utc()}", "",
             f"- instance rows: {len(inst)} (live {len(live)}; quarantined papers {len(q)})",
             f"- papers with rows: {len({r['paper_id'] for r in inst})}",
             f"- events (main + terms, origin/repeat): {len({(r['paper_id'], r['q_no'], r.get('side', '')) for r in ev})}",
             f"- classified rows: {sum(1 for r in live if code_of(r))}", "", "| year | exam | papers |", "|---|---|---|"]
    lines += [f"| {y} | {e} | {n} |" for (y, e), n in sorted(by_exam.items())]
    lines += [f"- {rel}: {len(read_table(v, rel)[1])} rows" for rel in SCHEMA]
    open(os.path.join(l2, "COUNTS.md"), "w").write("\n".join(lines) + "\n")
    # coverage + gaps
    cov = defaultdict(lambda: {"QP": "", "MS": ""})
    for f in files:
        if f["kind"] in ("QP", "MS", "SQP", "SQP-MS"):
            key = (f["year"], f["exam"], f["paper_code"] or "SQP")
            cov[key]["MS" if f["kind"] in ("MS", "SQP-MS") else "QP"] = f["status"]
    cov_rows = [{"year": k[0], "exam": k[1], "paper_code": k[2], "QP": d["QP"] or "missing", "MS": d["MS"] or "missing"}
                for k, d in sorted(cov.items())]
    write_csv(os.path.join(l2, "COVERAGE.csv"), cov_rows, ["year", "exam", "paper_code", "QP", "MS"])
    unread = table(v, "l1/UNREADABLE.csv")
    g = ["# GAPS (computed)", "", "## Question papers without a marking scheme"]
    g += [f"- {r['year']} {r['exam']} {r['paper_code']}" for r in cov_rows if r["QP"] not in ("missing",) and r["MS"] == "missing"] or ["- none"]
    g += ["", "## Files not yet mined"] + ([f"- {f['name']} ({f['year']} {f['kind']})" for f in files if f["status"] in ("NEW", "IN-PROGRESS")] or ["- none"])
    g += ["", "## Year not proven"] + ([f"- {f['name']}: {f['year_evidence'] or 'no evidence'}" for f in files if f["year"] == "UNK"] or ["- none"])
    g += ["", "## Unreadable content"] + ([f"- {r['paper_id']} p{r['page']} Q{r['q_no']}: {r['problem']} - {r['what_lost']}" for r in unread] or ["- none"])
    g += ["", "## Quarantined papers"] + ([f"- {p}" for p in sorted(q)] or ["- none"])
    open(os.path.join(l2, "GAPS.md"), "w").write("\n".join(g) + "\n")
    # statistics
    stats, target = archetype_stats(v, quarantine=quarantine)
    write_csv(os.path.join(l2, "ARCHETYPE-STATS.csv"), stats)
    st_by = {s["code"]: s for s in stats}
    # chapter marks per paper (side A only for OR)
    bps = {r["paper_id"]: r for r in table(v, "l1/BLUEPRINTS.csv")}
    ch2g = {r["chapter"]: r["mark_group"] for r in syl}
    cm = []
    for pid in sorted({r["paper_id"] for r in live if r["paper_id"] in bps}):
        seen, tot = set(), Counter()
        for r in live:
            if r["paper_id"] != pid or r.get("vi_alt") == "y" or r.get("side", "") == "B":
                continue
            m = fnum(r["marks_part"])
            if m is None:
                if r["q_no"] in seen:
                    continue
                m = fnum(r["q_marks_printed"]) or 0
                seen.add(r["q_no"])
            tot[r["chapter"]] += m
        cm += [{"paper_id": pid, "bp_class": bps[pid]["bp_class"], "chapter": ch, "mark_group": ch2g.get(ch, ""), "marks": f"{m:g}"}
               for ch, m in sorted(tot.items())]
    write_csv(os.path.join(l2, "CHAPTER-MARKS.csv"), cm, ["paper_id", "bp_class", "chapter", "mark_group", "marks"])
    # registers
    rg = os.path.join(l2, "registers")
    def reg_write(name, rows, cols):
        write_csv(os.path.join(rg, name + ".csv"), rows, cols)
        md = [f"# {name} (computed)", ""]
        for ch in sorted({r.get("chapter", "") for r in rows}):
            md += [f"## Chapter {ch} {chn.get(ch, '')}", "", md_table([r for r in rows if r.get("chapter", "") == ch], cols[:9])]
        open(os.path.join(rg, name + ".md"), "w").write("\n".join(md) + "\n")
    for name, pred in REGISTERS:
        cols = BASE_COLS + (["givens", "unknown"] if name.startswith("R02") else []) + (["options"] if name in ("R09-mcq", "R10-assertion-reason") else [])
        reg_write(name, [dict(r, code=code_of(r)) for r in live if pred(r)], cols)
    cases = table(v, "l1/CASES.csv")
    reg_write("R11-case-studies", cases, SCHEMA["l1/CASES.csv"])
    pairs = defaultdict(list)
    for r in live:
        if r.get("choice_pair_id"):
            pairs[r["choice_pair_id"]].append(r)
    r12 = []
    for k, rs in sorted(pairs.items()):
        sa = [r for r in rs if r.get("side") == "A"]
        sb = [r for r in rs if r.get("side") == "B"]
        r12.append({"choice_pair_id": k, "chapter": (sa or rs)[0]["chapter"],
                    "A_codes": " ".join(sorted({code_of(r) for r in sa if code_of(r)})), "B_codes": " ".join(sorted({code_of(r) for r in sb if code_of(r)})),
                    "A_chapters": " ".join(sorted({r["chapter"] for r in sa})), "B_chapters": " ".join(sorted({r["chapter"] for r in sb})),
                    "A_types": " ".join(sorted({r["q_type"] for r in sa})), "B_types": " ".join(sorted({r["q_type"] for r in sb}))})
    reg_write("R12-or-pairs", r12, ["choice_pair_id", "chapter", "A_codes", "B_codes", "A_chapters", "B_chapters", "A_types", "B_types"])
    ch_of = {r["instance_id"]: r["chapter"] for r in inst}
    reg_write("R13-formulas", [dict(r, chapter=ch_of.get(r["instance_id"], "")) for r in table(v, "l1/FORMULAE.csv")],
              ["chapter"] + SCHEMA["l1/FORMULAE.csv"])
    reg_write("R13-constants", [dict(r, chapter="") for r in table(v, "l1/CONSTANTS.csv")], ["chapter"] + SCHEMA["l1/CONSTANTS.csv"])
    conv = open(rpath(v, "CONVENTIONS.md")).read() if os.path.exists(rpath(v, "CONVENTIONS.md")) else ""
    r14 = [{"chapter": "", "rule": ln.strip("- ").strip()} for ln in conv.splitlines() if "examiner" in ln.lower() or "EXAMINER" in ln]
    reg_write("R14-examiner-rules", r14, ["chapter", "rule"])
    asked = {r["topic_id"] for r in ev if r["topic_id"]}
    r15 = [{"chapter": s["chapter"], "topic_id": s["topic_id"], "topic": s["topic"], "status": "never asked"}
           for s in syl if s["status"] == "IN" and s["topic_id"] not in asked]
    r15 += [{"chapter": r["chapter"], "topic_id": r["topic_id"], "topic": r["stem"][:80], "status": f"asked {r['year']} but {r['syllabus_status']} - skip"}
            for r in live if r["syllabus_status"] in ("DELETED", "VERIFY-DEL")]
    reg_write("R15-not-asked-and-deleted", r15, ["chapter", "topic_id", "topic", "status"])
    sq, mains = defaultdict(set), defaultdict(set)
    for r in live:
        c = code_of(r)
        if c and year_int(r):
            (sq if r["exam"] == "SQP" else mains if r["exam"] == "MAIN" else defaultdict(set))[year_int(r)].add(c)
    r16 = [{"chapter": "", "exam_year": y, "sqp_codes": len(sq[y]), "on_board_papers": len(sq[y] & mains[y]),
            "carryover": f"{len(sq[y] & mains[y]) / len(sq[y]):.2f}"} for y in sorted(sq) if mains.get(y)]
    reg_write("R16-sqp-carryover", r16, ["chapter", "exam_year", "sqp_codes", "on_board_papers", "carryover"])
    r17 = [dict(r, code=code_of(r), ncert_ref=cls[r["instance_id"]]["ncert_ref"]) for r in live
           if cls.get(r["instance_id"], {}).get("ncert_ref")]
    reg_write("R17-source-trace", r17, BASE_COLS + ["ncert_ref"])
    traps = [dict(r, code=code_of(r)) for r in live if "trap" in (r.get("note", "") + r.get("extra", "")).lower()]
    traps += [{"code": r["code"], "chapter": r["chapter"], "stem": r["traps"], "instance_id": "", "year": "", "exam": "",
               "paper_code": "", "q_no": "", "side": "", "part": "", "q_marks_printed": "", "q_type": r["q_type"],
               "command_verb": "", "rel": "", "note": "registry"} for r in reg.values() if r.get("traps") and r["status"] == "VERIFIED"]
    reg_write("R18-traps", traps, BASE_COLS + ["note"])
    # chapter cards — the archetype map
    cd = os.path.join(l2, "cards")
    os.makedirs(cd, exist_ok=True)
    steps_by = defaultdict(list)
    for s in steps:
        steps_by[s["instance_id"]].append(s)
    inst_by_code = defaultdict(list)
    for r in live:
        c = code_of(r)
        if c:
            inst_by_code[c].append(r)
    chapters = sorted({r["chapter"] for r in reg.values()} | {r["chapter"] for r in live if r["chapter"]})
    for ch in chapters:
        codes = [c for c, r in reg.items() if r["chapter"] == ch and r["status"] == "VERIFIED"]
        codes.sort(key=lambda c: -float(st_by.get(c, {}).get("score_uncalibrated", 0) or 0))
        md = [f"# Chapter {ch} · {chn.get(ch, '')} — archetype card (computed)", f"Stamp: pyqkit {KIT_EDITION} · {now_utc()}", ""]
        for c in codes:
            r = reg[c]
            s = st_by.get(c, {})
            md += [f"## {c} · {r['name']}", f"- **Asks you to:** {r['asks_you_to'].replace('Asks you to ', '', 1)}",
                   f"- **Type:** {r['q_type']} · **Family:** {r['family']} · **Topic:** {r['topic_id']}",
                   f"- **Seen:** {s.get('years', '') or 'only in non-main papers so far'} · events {s.get('events', 0)} · marks {s.get('marks_seen', '')} · score {s.get('score_uncalibrated', '')} (uncalibrated)",
                   f"- **Skeleton:** {r['skeleton'] or '(to write)'}", f"- **Traps:** {r['traps'] or '(to write)'}"]
            ex = sorted(inst_by_code[c], key=lambda x: x["year"], reverse=True)
            with_steps = next((x for x in ex if steps_by.get(x["instance_id"])), None)
            if with_steps:
                md.append(f"- **Scoring line** ({with_steps['instance_id']}):")
                md += [f"    - {st['marks']} · {st['step_text']} [{st['provenance']}]" for st in sorted(steps_by[with_steps["instance_id"]], key=lambda z: z["step_no"])]
            md += [f"- **Examples:** " + ", ".join(x["instance_id"] for x in ex[:5]), ""]
        drafts = [c for c, r in reg.items() if r["chapter"] == ch and r["status"] == "DRAFT"]
        if drafts:
            md += ["## Not yet seen in any mined paper (draft seeds)", ""] + [f"- {c} · {reg[c]['name']}" for c in sorted(drafts)]
        open(os.path.join(cd, f"CH{ch}.md"), "w").write("\n".join(md) + "\n")
    # year reports + index
    yd = os.path.join(l2, "years")
    os.makedirs(yd, exist_ok=True)
    for y in sorted({r["year"] for r in inst if r["year"]}):
        open(os.path.join(yd, f"Y{y}.md"), "w").write(year_report(v, y, inst, files, cls, code_of, steps, stats, cm, q))
    build_index(v, quiet=True)
    print(f"derive: l2 rebuilt · {len(stats)} archetypes scored for {target} · {len(cov_rows)} coverage cells · "
          f"{len(chapters)} chapter cards · quarantined {len(q)}")

def year_report(v, y, inst, files, cls, code_of, steps, stats, cm, quarantine):
    rows = [r for r in inst if r["year"] == y]
    yf = [f for f in files if f["year"] == y]
    papers = sorted({r["paper_id"] for r in rows})
    stepped = {s["instance_id"] for s in steps}
    prov = Counter(("TRANSFERRED" if s["provenance"].startswith("TRANSFERRED") else s["provenance"]) for s in steps
                   if s["instance_id"] in {r["instance_id"] for r in rows})
    new_codes = [s for s in stats if str(s["first_seen"]) == str(y)]
    reps = [r for r in rows if r.get("rel") == "repeat"]
    unread = [u for u in table(v, "l1/UNREADABLE.csv") if u["paper_id"] in papers]
    note_p = rpath(v, f"notes/Y{y}.md")
    out = [f"# YEAR REPORT {y} (computed)", f"pyqkit {KIT_EDITION} · {now_utc()}", "",
           f"- files: {len(yf)} · papers with rows: {len(papers)} · rows: {len(rows)} · classified: {sum(1 for r in rows if code_of(r))}",
           f"- scoring lines: " + (", ".join(f"{k} {n}" for k, n in prov.items()) or "none"),
           f"- archetypes first seen this year: {len(new_codes)} · repeats of older questions: {len(reps)}", "",
           "## Files", md_table(yf, ["file_id", "kind", "exam", "paper_code", "status", "name"]),
           "## Chapter marks per paper", md_table([c for c in cm if c["paper_id"] in papers], ["paper_id", "chapter", "mark_group", "marks"]),
           "## Archetypes first seen", md_table(new_codes, ["code", "name", "q_type", "chapter"]),
           "## Repeats", md_table(reps, ["instance_id", "rel_to", "q_type"]),
           "## Unreadable", md_table(unread, ["paper_id", "page", "q_no", "problem", "what_lost"]),
           "## Quarantined", "\n".join(f"- {p}" for p in papers if p in quarantine) or "- none", "",
           "## Observations (judged, from notes)", open(note_p).read() if os.path.exists(note_p) else "_none written_"]
    return "\n".join(out) + "\n"

def build_index(v, quiet=False):
    inst = table(v, "l1/INSTANCES.csv")
    canon, reg = canon_map(v)
    cls = {r["instance_id"]: canon(r["code"]) for r in table(v, "j/CLASSIFY.csv") if r["code"]}
    chn = chapter_names(v)
    st = load_state(v)
    out = ["# INDEX (computed)", f"pyqkit {KIT_EDITION} · {now_utc()} · batch {st.get('batch')} · years closed: {', '.join(map(str, st.get('years_closed', []))) or 'none'}", "",
           "## Vault files", md_table([{"file": rel, "rows": len(read_table(v, rel)[1])} for rel in SCHEMA], ["file", "rows"]),
           "## Papers by year"]
    for y in sorted({r["year"] for r in inst}, reverse=True):
        ps = sorted({(r["paper_id"], r["exam"], r["paper_code"]) for r in inst if r["year"] == y})
        out.append(f"- **{y}:** " + ", ".join(f"{p} ({e}, {c or '-'})" for p, e, c in ps))
    out += ["", "## Chapters -> archetypes -> questions"]
    by_code = defaultdict(list)
    for iid, c in cls.items():
        by_code[c].append(iid)
    for ch in sorted({r["chapter"] for r in reg.values()}):
        out.append(f"### Chapter {ch} {chn.get(ch, '')}")
        for c in sorted(k for k, r in reg.items() if r["chapter"] == ch and r["status"] in ("VERIFIED", "DRAFT")):
            ids = sorted(by_code.get(c, []))
            out.append(f"- {c} [{reg[c]['status']}] {reg[c]['name']} · {len(ids)} rows" + (": " + ", ".join(ids[:5]) + (" ..." if len(ids) > 5 else "") if ids else ""))
    regs = sorted(glob.glob(rpath(v, "l2/registers/*.csv")))
    out += ["", "## Registers"] + [f"- {os.path.basename(p)[:-4]}: {max(0, sum(1 for _ in open(p, encoding='utf-8')) - 1)} rows" for p in regs]
    prog = open(rpath(v, "PROGRESS.md")).read() if os.path.exists(rpath(v, "PROGRESS.md")) else ""
    open_items = [ln for ln in prog.split("## Open", 1)[-1].split("##", 1)[0].splitlines() if ln.strip().startswith("-")] if "## Open" in prog else []
    out += ["", "## Open items"] + (open_items or ["- none"])
    open(rpath(v, "l2/INDEX.md"), "w").write("\n".join(out) + "\n")
    if not quiet:
        print("index: l2/INDEX.md written")

def cmd_index(a):
    build_index(a.vault)

# ------------------------------------------------------------------------------------------------
# year-close — the Definition of Done (CI§5)
# ------------------------------------------------------------------------------------------------
def cmd_year_close(a):
    v, y = a.vault, str(a.year)
    inst = table(v, "l1/INSTANCES.csv")
    files = table(v, "l1/FILES.csv")
    rows = [r for r in inst if r["year"] == y and r.get("vi_alt") != "y"]
    byid = {r["instance_id"]: r for r in inst}
    canon, reg = canon_map(v)
    cls = {r["instance_id"]: canon(r["code"]) for r in table(v, "j/CLASSIFY.csv") if r["code"]}
    steps_ids = {s["instance_id"] for s in table(v, "l1/STEP-AWARDS.csv")}
    figs = {f["instance_id"] for f in table(v, "l1/FIGURES.csv")}
    forms = {f["instance_id"] for f in table(v, "l1/FORMULAE.csv")}
    cases = {c["case_id"] for c in table(v, "l1/CASES.csv")}
    consts = {c["paper_id"] for c in table(v, "l1/CONSTANTS.csv")}
    bps = {r["paper_id"]: r for r in table(v, "l1/BLUEPRINTS.csv")}
    unread = {(u["paper_id"], u["q_no"]) for u in table(v, "l1/UNREADABLE.csv")}
    origin = lambda r: byid.get(r.get("rel_to", "")) if r.get("rel") in ("shuffle", "variant") else None
    def has(idset, r):
        o = origin(r)
        return r["instance_id"] in idset or (o is not None and o["instance_id"] in idset)
    checks = []
    def add(n, name, fails):
        checks.append((n, name, "PASS" if not fails else "FAIL", "; ".join(fails[:6]) + (f" (+{len(fails) - 6} more)" if len(fails) > 6 else "")))
    yfiles = [f for f in files if f["year"] == y and f["kind"] not in ("OTHER", "SYLLABUS")]
    add(1, "every file of the year DONE / DUP-FILE / UNAVAILABLE", [f"{f['name']} is {f['status']}" for f in yfiles if f["status"] not in ("DONE", "DUP-FILE", "UNAVAILABLE", "QUARANTINE")])
    qp = sorted({f["paper_id"] for f in yfiles if f["kind"] in ("QP", "SQP", "APQ") and f["paper_id"]})
    f2 = []
    for pid in qp:
        st = next((f["status"] for f in yfiles if f["paper_id"] == pid and f["kind"] in ("QP", "SQP", "APQ")), "")
        if pid not in bps:
            f2.append(f"{pid}: no blueprint")
            continue
        probs, _ = check_paper(v, pid, inst, bps)
        if probs and st != "QUARANTINE":
            f2.append(f"{pid}: {probs[0]}")
    add(2, "blueprint + check-paper for every question paper", f2)
    add(3, "every question paper has its rows", [pid for pid in qp if not any(r["paper_id"] == pid for r in rows)])
    f4 = [f"{r['instance_id']}: {c}" for r in rows for c in ("chapter", "topic_id", "q_type", "command_verb")
          if (not r[c] or r[c] == "UNK") and (r["paper_id"], r["q_no"]) not in unread]
    add(4, "chapter, topic, type, verb on every row", f4)
    add(5, "every row classified to a VERIFIED code", [r["instance_id"] for r in rows if reg.get(cls.get(r["instance_id"], ""), {}).get("status") != "VERIFIED"])
    qs = defaultdict(list)
    for r in rows:
        qs[(r["paper_id"], r["q_no"], r.get("side", ""))].append(r)
    no_steps = [f"{p} Q{qn}{s}" for (p, qn, s), rs in qs.items() if not any(has(steps_ids, r) for r in rs)]
    add(6, f"scoring lines for every question ({len(qs) - len(no_steps)}/{len(qs)})", no_steps)
    f7 = [f"{r['instance_id']}: figure" for r in rows if r.get("has_figure") == "y" and not has(figs, r)]
    f7 += [f"{r['instance_id']}: formula" for r in rows if r["q_type"] in ("NUM", "DRV") and not has(forms, r)]
    f7 += [f"{r['case_id']}: case" for r in rows if r.get("case_id") and r["case_id"] not in cases]
    f7 += [f"{pid}: constants" for pid in qp if bps.get(pid, {}).get("constants_listed") == "y" and pid not in consts]
    add(7, "figures, formulas, cases, constants complete", sorted(set(f7)))
    used = sorted({cls[r["instance_id"]] for r in rows if r["instance_id"] in cls})
    add(8, "every code used has name, asks_you_to, skeleton, traps",
        [f"{c}: {k}" for c in used for k in ("name", "asks_you_to", "skeleton", "traps") if not reg.get(c, {}).get(k, "").strip()])
    need = [rpath(v, "l2/INDEX.md"), rpath(v, f"l2/years/Y{y}.md")] + [rpath(v, f"l2/registers/{n}.md") for n, _ in REGISTERS]
    need += [rpath(v, f"l2/cards/CH{ch}.md") for ch in sorted({r["chapter"] for r in rows if r["chapter"] and r["chapter"] != "X11"})]
    xl = os.path.join(a.xlsx_dir, f"PHY-YEAR-{y}.xlsx")
    add(9, "registers, cards, index, year report, workbook generated", [os.path.basename(p) for p in need + [xl] if not os.path.exists(p)])
    res, _ = validate(v)
    add(10, "all validation checks pass", [f"{c} {n}" for c, n, r, d in res if r == "FAIL"])
    ok = all(c[2] == "PASS" for c in checks)
    print(f"# YEAR-CLOSE {y}: {'PASS' if ok else 'NOT YET'}\n\n| # | Definition of Done | Result | Left to do |\n|---|---|---|---|")
    for n, name, r, d in checks:
        print(f"| {n} | {name} | {r} | {d.replace('|', '/')} |")
    if ok:
        st = load_state(v)
        closed = [str(x) for x in st.get("years_closed", [])]
        if y not in closed:
            closed.append(y)
        st["years_closed"], st["year_in_progress"] = closed, ""
        save_state(v, st)
        print(f"\nYear {y} closed. Your three steps: (1) download the new PHY-VAULT file; (2) replace the old vault file in "
              f"project knowledge; (3) swap the marking schemes in project knowledge to the next year.")
    sys.exit(0 if ok else 1)

# ------------------------------------------------------------------------------------------------
# export-xlsx — the year workbook (xlsx skill rules: Arial, formulas, recalc)
# ------------------------------------------------------------------------------------------------
def cmd_export_xlsx(a):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    v, y = a.vault, str(a.year)
    inst = [r for r in table(v, "l1/INSTANCES.csv") if r["year"] == y]
    ids = {r["instance_id"] for r in inst}
    canon, reg = canon_map(v)
    cls = {r["instance_id"]: canon(r["code"]) for r in table(v, "j/CLASSIFY.csv") if r["code"]}
    bps = {r["paper_id"]: r for r in table(v, "l1/BLUEPRINTS.csv")}
    files = [f for f in table(v, "l1/FILES.csv") if f["year"] == y]
    wb = Workbook()
    def sheet(title, cols, rows, first=False):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title = title
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols, 1):
            width = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
        return ws
    qcols = ["instance_id", "paper_id", "q_no", "side", "part", "section", "marks_part", "q_marks_printed", "chapter",
             "topic_id", "q_type", "command_verb", "code", "stem", "rel", "source_page"]
    papers = []
    for pid in sorted({r["paper_id"] for r in inst}):
        probs, _ = check_paper(v, pid)
        f = next((x for x in files if x["paper_id"] == pid), {})
        papers.append({"paper_id": pid, "exam": f.get("exam", ""), "paper_code": f.get("paper_code", ""), "file": f.get("name", ""),
                       "status": f.get("status", ""), "questions": len({r["q_no"] for r in inst if r["paper_id"] == pid}),
                       "check": "PASS" if not probs else "QUARANTINE", "bp_class": bps.get(pid, {}).get("bp_class", "")})
    ws = sheet("Summary", ["Measure", "Value"], [], first=True)
    sheet("Papers", ["paper_id", "exam", "paper_code", "file", "status", "questions", "check", "bp_class"], papers)
    sheet("Questions", qcols, [dict(r, code=cls.get(r["instance_id"], "")) for r in inst])
    sheet("Scoring", ["award_id", "instance_id", "step_no", "step_text", "marks", "step_kind", "provenance"],
          [s for s in table(v, "l1/STEP-AWARDS.csv") if s["instance_id"] in ids])
    used = Counter(cls[i] for i in ids if i in cls)
    sheet("Archetypes", ["code", "name", "asks_you_to", "chapter", "q_type", "rows_this_year"],
          [dict(reg.get(c, {}), code=c, rows_this_year=n) for c, n in sorted(used.items())])
    sheet("Figures", ["figure_id", "instance_id", "fig_type", "printed_or_demanded", "description"],
          [f for f in table(v, "l1/FIGURES.csv") if f["instance_id"] in ids])
    sheet("Formulas", ["formula_id", "instance_id", "role", "formula", "units"], [f for f in table(v, "l1/FORMULAE.csv") if f["instance_id"] in ids])
    qt = get_column_letter(qcols.index("q_type") + 1)
    summary = [("Year", y), ("Question rows", "=COUNTA(Questions!A:A)-1"), ("Papers", "=COUNTA(Papers!A:A)-1"),
               ("Scoring lines", "=COUNTA(Scoring!A:A)-1"), ("Archetypes used", "=COUNTA(Archetypes!A:A)-1"),
               ("MCQ rows", f'=COUNTIF(Questions!{qt}:{qt},"MCQ")'), ("Numerical rows", f'=COUNTIF(Questions!{qt}:{qt},"NUM")'),
               ("Derivation rows", f'=COUNTIF(Questions!{qt}:{qt},"DRV")'),
               ("Note", "Counts are live formulas over the sheets of this workbook (generated by pyqkit).")]
    for m, val in summary:
        ws.append([m, val])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 70
    os.makedirs(a.out, exist_ok=True)
    path = os.path.join(a.out, f"PHY-YEAR-{y}.xlsx")
    wb.save(path)
    rc = "/mnt/skills/public/xlsx/scripts/recalc.py"
    note = ""
    if os.path.exists(rc):
        r = subprocess.run([sys.executable, rc, path, "90"], capture_output=True, text=True)
        try:
            j = json.loads(r.stdout)
            note = f" · recalc {j.get('status')}: {j.get('total_formulas', 0)} formulas, {j.get('total_errors', 0)} errors"
        except (json.JSONDecodeError, ValueError):
            note = " · recalc output unreadable: " + (r.stdout or r.stderr).strip()[-120:]
    print(f"export-xlsx: {path} · {len(inst)} question rows{note}")

# ------------------------------------------------------------------------------------------------
# plan-mock — blueprint-exact slot plan (P§15.3)
# ------------------------------------------------------------------------------------------------
DEFAULT_SECTIONS = "A:1-16:1;B:17-21:2;C:22-28:3;D:29-30:4;E:31-33:5"
DEFAULT_CHOICE = {20, 21, 27, 31, 32, 33}
DEFAULT_TYPES = {"A-MCQ": {"MCQ"}, "A-AR": {"AR"}, "B": {"CON", "NUM", "DEF", "PRD", "CMP", "DGM", "GRF"},
                 "C": {"NUM", "DRV", "CON", "DGM", "GRF", "DEV", "PRD", "CMP"}, "D": {"CON", "NUM", "PRD", "GRF", "DEF", "MCQ"},
                 "E": {"DRV", "NUM", "DEV", "GRF"}}
AS_MCQ = {"DEF", "CON", "PRD", "NUM", "GRF", "CMP"}
COG_TARGET = {"REM+UND": 27, "APP": 22, "ANL": 21}

def parse_choice(rule):
    qs = set()
    for part in (rule or "").split(";"):
        if ":" in part:
            qs |= {int(x) for x in re.findall(r"\d+", part.split(":", 1)[1])}
    return qs

def _plan_once(a):
    v = a.vault
    rng = random.Random(a.seed)
    bps = {r["paper_id"]: r for r in table(v, "l1/BLUEPRINTS.csv")}
    bp = bps.get("SQP2627-042", {})
    secmap = parse_sections(bp.get("sections") or DEFAULT_SECTIONS)
    choice = parse_choice(bp.get("choice_rule")) or DEFAULT_CHOICE
    syl = table(v, "j/SYLLABUS.csv")
    ch2g = {r["chapter"]: r["mark_group"] for r in syl}
    quota = {}
    for r in syl:
        if r["mark_group"].startswith("G"):
            quota[r["mark_group"]] = int(float(r["group_marks"] or 0))
    topic = {r["topic_id"]: r for r in syl}
    canon, reg = canon_map(v)
    _, quarantine = validate(v)
    stats, target = archetype_stats(v, quarantine=quarantine)
    st_by = {s["code"]: s for s in stats}
    cls_rows = table(v, "j/CLASSIFY.csv")
    cog = defaultdict(Counter)
    for r in cls_rows:
        if r["code"]:
            cog[canon(r["code"])][r["cog_level"]] += 1
    used_before = Counter()
    for m in table(v, "out/MOCKS.csv"):
        used_before.update(c for c in m["codes"].split() if c)
    inst = table(v, "l1/INSTANCES.csv")
    struct = Counter((r["section"], r["q_type"]) for r in inst if bps.get(r["paper_id"], {}).get("bp_class") == "P33")
    allowed = {k: set(t) for k, t in DEFAULT_TYPES.items()}
    for sec in ("B", "C", "E"):
        seen = {t for (s, t), n in struct.items() if s == sec}
        if len(seen) >= 3:
            allowed[sec] = seen
    def blocked(r):
        if r["status"] != "VERIFIED":
            return True
        t = topic.get(r["topic_id"])
        if t and t["status"] != "IN":
            return True
        lim = (t or {}).get("limitation", "")
        return (r["q_type"] == "DRV" and "no derivation" in lim) or (r["q_type"] == "NUM" and "qualitative" in lim)
    pool = [r for r in reg.values() if not blocked(r) and (not a.chapter or r["chapter"] == a.chapter)]
    if not pool:
        sys.exit("plan-mock: no VERIFIED archetypes available yet - mine and close at least one year first")
    def sc(code):
        return float(st_by.get(code, {}).get("score_uncalibrated", 0) or 0)
    def key(r):
        s = st_by.get(r["code"], {})
        base = sc(r["code"])
        if a.mode == "stress":
            return (r["q_type"] in ("NUM", "DRV"), -float(s.get("p_year") or 0), base)
        if a.mode == "wildcard":
            return (base + (0.3 if int(s.get("n_years") or 0) <= 1 else 0),)
        if a.mode == "coverage":
            return (base - 0.5 * used_before[r["code"]],)
        return (base,)
    def pick(cands):
        if not cands:
            return None
        if a.mode == "seed":
            weights = [sc(r["code"]) + 0.05 for r in cands]
            return rng.choices(cands, weights=weights, k=1)[0]
        best = max(key(r) for r in cands)
        top = [r for r in cands if key(r) == best]
        return rng.choice(top)
    remaining = dict(quota)
    used, fam_n, plan = set(), Counter(), []
    cog_now = Counter()
    def lvl_of(code):
        lv = cog[code].most_common(1)[0][0] if cog[code] else ""
        return "REM+UND" if lv in ("REM", "UND") else lv
    def steer(r):
        lv = lvl_of(r["code"])
        if a.mode == "chapter" or lv not in COG_TARGET:
            return 0.0
        return 0.15 * max(0.0, COG_TARGET[lv] - cog_now[lv]) / COG_TARGET[lv]
    def deficit(g):
        """marks still needed in group g minus unused archetypes able to take them (scarcity signal)"""
        return remaining.get(g, 0) - sum(1 for r in pool if r["code"] not in used and ch2g.get(r["chapter"], "") == g)
    def key2(r):
        k0 = key(r)
        k0 = (k0[0] + steer(r),) + tuple(k0[1:]) if len(k0) == 1 else k0
        return (1 if deficit(ch2g.get(r["chapter"], "")) > 0 else 0,) + tuple(k0)
    def pick2(cands):
        if not cands:
            return None
        if a.mode == "seed":
            return rng.choices(cands, weights=[(sc(r["code"]) + 0.05 + steer(r)) * (5 if deficit(ch2g.get(r["chapter"], "")) > 0 else 1)
                                               for r in cands], k=1)[0]
        best = max(key2(r) for r in cands)
        return rng.choice([r for r in cands if key2(r) == best])
    order = sorted(secmap.items(), key=lambda kv: ("EDCBA".index(kv[1][0]) if kv[1][0] in "EDCBA" else 9, kv[0]))
    for qn, (sec, marks) in order:
        m = int(marks)
        slot = "A-AR" if sec == "A" and qn >= 13 else ("A-MCQ" if sec == "A" else sec)
        types = allowed.get(slot, set())
        def ok(r, group=None, need_quota=True, type_set=None, fam_cap=2):
            g = ch2g.get(r["chapter"], "")
            if r["code"] in used or fam_n[fam_key(r["family"])] >= fam_cap:
                return False
            if group and g != group:
                return False
            if need_quota and not a.chapter and remaining.get(g, 0) < m:
                return False
            return r["q_type"] in (type_set if type_set is not None else types)
        groups = [None]
        if sec == "A" and not a.chapter:     # 1-mark slots fill the group under most pressure first
            groups = sorted([g for g in remaining if remaining[g] >= m], key=lambda g: (-deficit(g), -remaining[g])) or [None]
        c, as_type, flag = None, "", ""
        mcq_types = AS_MCQ if slot in ("A-MCQ", "A-AR") else None
        for grp in groups:
            steps_ = [(dict(group=grp), "", ""),
                      (dict(group=grp, type_set=mcq_types), ("as MCQ" if slot == "A-MCQ" else "as A-R"), "") if mcq_types else None,
                      (dict(group=grp, type_set=ENUM["q_type"]), "", "type relaxed"),
                      (dict(group=grp, type_set=ENUM["q_type"], fam_cap=99), "", "family cap relaxed")]
            for st in [x for x in steps_ if x]:
                c = pick2([r for r in pool if ok(r, **st[0])])
                if c is not None:
                    as_type, flag = st[1], st[2]
                    break
            if c is not None:
                break
        if c is None:
            c = pick2([r for r in pool if ok(r, need_quota=False, type_set=ENUM["q_type"], fam_cap=99)])
            flag = "quota overflow"
        if c is None:
            plan.append({"q_no": qn, "side": "", "section": sec, "marks": m, "code": "", "flag": "no candidate"})
            continue
        g = ch2g.get(c["chapter"], "")
        remaining[g] = remaining.get(g, 0) - m
        used.add(c["code"])
        fam_n[fam_key(c["family"])] += 1
        cog_now[lvl_of(c["code"])] += m
        s_ = st_by.get(c["code"], {})
        plan.append({"q_no": qn, "side": "A" if qn in choice else "", "section": sec, "marks": m, "code": c["code"], "name": c["name"],
                     "chapter": c["chapter"], "group": g, "q_type": c["q_type"], "as_type": as_type,
                     "cog_level": (cog[c["code"]].most_common(1)[0][0] if cog[c["code"]] else ""),
                     "score": f"{sc(c['code']):.3f}", "years": s_.get("years", ""), "flag": flag})
        if qn in choice:
            alts = [r for r in pool if r["code"] not in used and ch2g.get(r["chapter"], "") == g
                    and fam_key(r["family"]) != fam_key(c["family"]) and r["q_type"] in (types or ENUM["q_type"])]
            b = pick2(alts) or pick2([r for r in pool if r["code"] not in used and ch2g.get(r["chapter"], "") == g])
            if b:
                used.add(b["code"])
                fam_n[fam_key(b["family"])] += 1
                plan.append({"q_no": qn, "side": "B", "section": sec, "marks": m, "code": b["code"], "name": b["name"],
                             "chapter": b["chapter"], "group": g, "q_type": b["q_type"], "as_type": "",
                             "cog_level": (cog[b["code"]].most_common(1)[0][0] if cog[b["code"]] else ""),
                             "score": f"{sc(b['code']):.3f}", "years": st_by.get(b["code"], {}).get("years", ""), "flag": ""})
            else:
                plan.append({"q_no": qn, "side": "B", "section": sec, "marks": m, "code": "", "flag": "no OR alternative"})
    plan.sort(key=lambda r: (r["q_no"], r["side"]))
    used_marks = Counter()
    cog_marks = Counter()
    for r in plan:
        if r.get("code") and r["side"] in ("", "A"):
            used_marks[r["group"]] += r["marks"]
            lv = r.get("cog_level", "")
            cog_marks["REM+UND" if lv in ("REM", "UND") else lv or "unknown"] += r["marks"]
    FATAL = {"quota overflow", "no candidate", "no OR alternative"}
    if a.chapter:                      # a chapter drill is as long as the chapter's distinct archetypes allow
        plan = [r for r in plan if r.get("code")]
        gate_ok = len(plan) >= min(8, len(pool))
    else:
        gate_ok = all(used_marks.get(g, 0) == quota[g] for g in quota) and not any(r.get("flag") in FATAL for r in plan)
    if not gate_ok and not getattr(a, "_final", True):
        return False
    out = a.out or rpath(v, "out/plans")
    os.makedirs(out, exist_ok=True)
    stem = f"PLAN-{a.mode}-{a.seed}-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}"
    cols = ["q_no", "side", "section", "marks", "code", "name", "chapter", "group", "q_type", "as_type", "cog_level", "score", "years", "flag"]
    write_csv(os.path.join(out, stem + ".csv"), plan, cols)
    quota_lines = [f"| {g} | {quota[g]} | {used_marks.get(g, 0)} | {'OK' if used_marks.get(g, 0) == quota[g] else 'CHECK'} |" for g in sorted(quota)]
    cog_lines = [f"| {k} | {COG_TARGET[k]} | {cog_marks.get(k, 0)} |" for k in COG_TARGET] + [f"| unknown | - | {cog_marks.get('unknown', 0)} |"]
    md = [f"# MOCK PLAN · mode {a.mode} · seed {a.seed} · target {target}", f"pyqkit {KIT_EDITION} · {now_utc()} · scores uncalibrated until backtested", "",
          md_table(plan, cols), "## Mark-group quotas (side A of each choice counts)", "", "| group | quota | planned | check |", "|---|---|---|---|",
          *quota_lines, "", "## Thinking-level mix (target 27 / 22 / 21)", "", "| level | target | planned |", "|---|---|---|", *cog_lines,
          "", f"**Quality gate:** {'PASS' if gate_ok else 'FAIL - quotas or flags; not to be shown as a mock'}"]
    open(os.path.join(out, stem + ".md"), "w").write("\n".join(md) + "\n")
    flags = [r for r in plan if r.get("flag")]
    print(f"plan-mock: {len([r for r in plan if r.get('code')])} slots filled · " +
          ("chapter drill (quotas not applied)" if a.chapter else "quotas " + ", ".join(f"{g} {used_marks.get(g, 0)}/{quota[g]}" for g in sorted(quota))) +
          f" · flags {len(flags)} · seed {a.seed} · QUALITY GATE {'PASS' if gate_ok else 'FAIL - do not show this plan'}"
          f" -> {os.path.join(out, stem)}.csv/.md")
    return gate_ok

def cmd_plan_mock(a):
    tries = max(1, a.tries)
    for i in range(tries):
        b = argparse.Namespace(**vars(a))
        b.seed, b._final = a.seed + i, (i == tries - 1)
        if _plan_once(b):
            return
    sys.exit(2)

# ------------------------------------------------------------------------------------------------
# backtest (P§16)
# ------------------------------------------------------------------------------------------------
def question_hits(v, year, codes):
    canon, _ = canon_map(v)
    cls = {r["instance_id"]: canon(r["code"]) for r in table(v, "j/CLASSIFY.csv") if r["code"]}
    qs = defaultdict(lambda: {"codes": set(), "marks": 0.0})
    for r in table(v, "l1/INSTANCES.csv"):
        if r["exam"] != "MAIN" or year_int(r) != year or r.get("vi_alt") == "y" or r.get("rel", "") not in EVENT_RELS:
            continue
        k = (r["paper_id"], r["q_no"], r.get("side", ""))
        if r["instance_id"] in cls:
            qs[k]["codes"].add(cls[r["instance_id"]])
        m = fnum(r["q_marks_printed"])
        if m is not None:
            qs[k]["marks"] = m
    qs = {k: d for k, d in qs.items() if d["codes"]}
    if not qs:
        return None
    hit = [k for k, d in qs.items() if d["codes"] & codes]
    tm = sum(d["marks"] for d in qs.values()) or 1
    return {"events": len(qs), "event_hit": len(hit) / len(qs), "marks_hit": sum(qs[k]["marks"] for k in hit) / tm}

def cmd_backtest(a):
    v, y = a.vault, int(a.year)
    canon, _ = canon_map(v)
    pred = set()
    with open(a.pred, newline="") as f:
        pred = {canon(r["code"].strip()) for r in csv.DictReader(f) if r.get("code")}
    cls = {r["instance_id"]: canon(r["code"]) for r in table(v, "j/CLASSIFY.csv") if r["code"]}
    by_year, sqp = defaultdict(set), set()
    for r in table(v, "l1/INSTANCES.csv"):
        c = cls.get(r["instance_id"])
        if not c or not year_int(r):
            continue
        if r["exam"] == "MAIN":
            by_year[year_int(r)].add(c)
        if r["exam"] == "SQP" and year_int(r) == y:
            sqp.add(c)
    pool = sorted(set().union(*[s for yy, s in by_year.items() if yy < y])) if any(yy < y for yy in by_year) else []
    lines = [f"# BACKTEST - predicting {y} (computed)", "", "| Prediction | codes | event hit | marks hit |", "|---|---|---|---|"]
    for name, s in (("engine", pred), ("SQP only", sqp), ("last year only", by_year.get(y - 1, set()))):
        m = question_hits(v, y, s) if s else None
        lines.append(f"| {name} | {len(s)} | " + (f"{m['event_hit']:.0%} | {m['marks_hit']:.0%} |" if m else "n/a | n/a |"))
    if pool and pred:
        r_ = random.Random(42)
        k = min(len(pred), len(pool))
        sims = [s for s in (question_hits(v, y, set(r_.sample(pool, k))) for _ in range(100)) if s]
        if sims:
            lines.append(f"| random x100 (mean) | {k} | {statistics.mean(s['event_hit'] for s in sims):.0%} | {statistics.mean(s['marks_hit'] for s in sims):.0%} |")
    base = question_hits(v, y, pred)
    lines += ["", f"Real {y} events scored: {base['events'] if base else 0}. Trust is claimed only where the engine beats every baseline."]
    os.makedirs(rpath(v, "l2"), exist_ok=True)
    open(rpath(v, "l2/BACKTEST.md"), "a").write("\n".join(lines) + "\n\n")
    print("\n".join(lines))

# ------------------------------------------------------------------------------------------------
# vault file (JSON) — pack / unpack / restore (P§3.2)
# ------------------------------------------------------------------------------------------------
def vault_files(v):
    out = []
    for f in sorted(glob.glob(os.path.join(v, "**", "*"), recursive=True)):
        rel = os.path.relpath(f, v)
        if os.path.isfile(f) and not rel.startswith(("work/", "work" + os.sep, "l2/", "l2" + os.sep)) and not rel.endswith(".tmp"):
            out.append((f, rel.replace(os.sep, "/")))
    return out

def total_rows(counts):
    return sum((counts or {}).values())

def cmd_pack(a):
    v = a.vault
    st = load_state(v)
    if a.batch is not None:
        st["batch"] = int(a.batch)
    st.update({"schema_edition": SCHEMA_EDITION, "kit_edition": KIT_EDITION, "created_utc": now_utc(), "rows": row_counts(v)})
    save_state(v, st)
    os.makedirs(rpath(v, "tools"), exist_ok=True)
    if os.path.abspath(__file__) != os.path.abspath(rpath(v, "tools/pyqkit.py")):
        shutil.copy2(os.path.abspath(__file__), rpath(v, "tools/pyqkit.py"))
    files = {}
    for f, rel in vault_files(v):
        raw = open(f, "rb").read()
        files[rel] = {"sha256": sha256_bytes(raw), "text": raw.decode("utf-8", "replace")}
    rows = total_rows(st["rows"])
    name = f"PHY-VAULT-B{int(st.get('batch', 0)):02d}-{dt.datetime.now().strftime('%Y%m%d-%H%M')}-{rows}r.json"
    os.makedirs(a.out, exist_ok=True)
    out = os.path.join(a.out, name)
    json.dump({"kind": "pyqkit-vault", "kit_edition": KIT_EDITION, "state": st, "files": files}, open(out, "w", encoding="utf-8"), ensure_ascii=False)
    print(out)

def load_candidate(path):
    """-> (state, files{rel: bytes}, bad_count) or (None, None, reason)."""
    try:
        if path.endswith(".json"):
            obj = json.load(open(path, encoding="utf-8"))
            if obj.get("kind") != "pyqkit-vault":
                return None, None, "not a vault file"
            files, bad = {}, 0
            for rel, meta in obj["files"].items():
                data = meta["text"].encode("utf-8")
                if sha256_bytes(data) != meta["sha256"]:
                    bad += 1
                files[rel] = data
            return obj.get("state", {}), files, bad
        with zipfile.ZipFile(path) as z:        # legacy ZIP vaults (PHY-VAULT/ prefix + CHECKSUMS)
            names = z.namelist()
            pre = "PHY-VAULT/" if any(n.startswith("PHY-VAULT/") for n in names) else ""
            st = json.loads(z.read(pre + "STATE.json"))
            files, bad = {}, 0
            sums = z.read(pre + "CHECKSUMS.sha256").decode().splitlines() if pre + "CHECKSUMS.sha256" in names else []
            want = {l.split("  ", 1)[1]: l.split("  ", 1)[0] for l in sums if "  " in l}
            for n in names:
                if n.endswith("/") or n.endswith("CHECKSUMS.sha256"):
                    continue
                rel = n[len(pre):]
                data = z.read(n)
                if rel in want and sha256_bytes(data) != want[rel]:
                    bad += 1
                files[rel] = data
            return st, files, bad
    except Exception as e:
        return None, None, str(e)[:80]

def install_files(files, vault):
    tmp = vault.rstrip("/") + "_restore_tmp"
    shutil.rmtree(tmp, ignore_errors=True)
    for rel, data in files.items():
        p = os.path.join(tmp, rel)
        os.makedirs(os.path.dirname(p), exist_ok=True)
        open(p, "wb").write(data)
    keep_work = os.path.join(vault, "work")
    if os.path.isdir(keep_work):
        shutil.move(keep_work, os.path.join(tmp, "work"))
    shutil.rmtree(vault, ignore_errors=True)
    shutil.move(tmp, vault)
    for sub in ("l2", "work", "notes", "out"):
        os.makedirs(os.path.join(vault, sub), exist_ok=True)

def cmd_unpack(a):
    st, files, bad = load_candidate(a.file)
    if st is None:
        sys.exit(f"unpack: {a.file} is not usable ({bad})")
    install_files(files, a.to)
    print(f"unpacked {os.path.basename(a.file)} -> {a.to} · batch {st.get('batch')} · rows {total_rows(st.get('rows'))} · checksum failures {bad}")

def cmd_restore(a):
    sources = a.sources or ["/mnt/project", "/mnt/user-data/uploads"]
    paths = []
    for s in sources:
        if os.path.isfile(s):
            paths.append(s)
        else:
            paths += glob.glob(os.path.join(s, "**", "PHY-VAULT*.json"), recursive=True)
            paths += glob.glob(os.path.join(s, "**", "PHY042-VAULT*.zip"), recursive=True)
    cands = []
    for p in sorted(set(paths)):
        st, files, bad = load_candidate(p)
        if st is None:
            print(f"  skip {os.path.basename(p)}: {bad}")
            continue
        cands.append((p, st, files, bad))
    if a.db and os.path.isdir(a.db):
        tmp = tempfile.mkdtemp()
        files, man, problems = db_collect(a.db)
        if man and not problems:
            cands.append((f"auto-vault ({a.db})", {"batch": man.get("batch", 0), "created_utc": man.get("created_utc", ""),
                          "rows": man.get("rows", {})}, files, 0))
        elif man:
            print(f"  auto-vault dump has problems: {'; '.join(problems[:3])}")
    if os.path.exists(rpath(a.vault, "STATE.json")):
        cur = {rel: open(f, "rb").read() for f, rel in vault_files(a.vault)}
        cands.append(("current working copy", load_state(a.vault), cur, 0))
    if not cands:
        print("restore: no vault found anywhere - start with `init` (Batch 0) or a recovery vault")
        sys.exit(3)
    def rank(c):
        p, st, files, bad = c
        rows = total_rows(st.get("rows"))
        return (bad == 0, rows > 0, int(st.get("batch", 0) or 0), str(st.get("created_utc", "")), rows)
    cands.sort(key=rank, reverse=True)
    for i, (p, st, files, bad) in enumerate(cands):
        verdict = "CHOSEN" if i == 0 else "stale"
        print(f"  {verdict:7} {os.path.basename(p) if os.path.isfile(p) else p} · batch {st.get('batch')} · rows {total_rows(st.get('rows'))} · "
              f"created {st.get('created_utc', '')} · checksum failures {bad}")
    p, st, files, bad = cands[0]
    if p != "current working copy":
        install_files(files, a.vault)
    if total_rows(st.get("rows")) == 0:
        print("  note: the chosen vault has 0 rows (Batch 0 state)")
    print(f"restore: vault ready at {a.vault} · batch {st.get('batch')} · rows {total_rows(st.get('rows'))}")

# ------------------------------------------------------------------------------------------------
# auto-vault (artifact database) — db-export / db-import
# ------------------------------------------------------------------------------------------------
DOC_LIMIT = 180_000

def cmd_db_export(a):
    v = a.vault
    st = load_state(v)
    shutil.rmtree(a.out, ignore_errors=True)
    os.makedirs(os.path.join(a.out, "docs"), exist_ok=True)
    manifest = {"kind": "pyqkit-manifest", "batch": st.get("batch", 0), "created_utc": now_utc(), "rows": row_counts(v),
                "kit_edition": KIT_EDITION, "files": {}}
    entries = []
    for f, rel in vault_files(v):
        raw = open(f, "rb").read()
        text = raw.decode("utf-8", "replace")
        chunks, cur = [], ""
        for line in text.splitlines(keepends=True):
            if len(cur.encode()) + len(line.encode()) > DOC_LIMIT and cur:
                chunks.append(cur)
                cur = ""
            cur += line
        chunks.append(cur)
        base = re.sub(r"[^A-Za-z0-9_.~:@+-]", "_", rel.replace("/", "__"))
        ids = []
        for i, ch in enumerate(chunks):
            did = f"{base}__{i:03d}"
            fp = os.path.join(a.out, "docs", did + ".json")
            json.dump({"kind": "pyqkit-doc", "path": rel, "part": i, "of": len(chunks), "text": ch}, open(fp, "w"), ensure_ascii=False)
            ids.append(did)
            entries.append({"op": "set", "collection": "vault", "doc_id": did, "file_path": fp})
        manifest["files"][rel] = {"sha256": sha256_bytes(raw), "bytes": len(raw), "docs": ids}
    mp = os.path.join(a.out, "docs", "manifest.json")
    json.dump(manifest, open(mp, "w"))
    entries.append({"op": "set", "collection": "meta", "doc_id": "manifest", "file_path": mp})
    plans = [entries[i:i + 50] for i in range(0, len(entries), 50)]
    for k, plan in enumerate(plans, 1):
        json.dump(plan, open(os.path.join(a.out, f"writes-{k:03d}.json"), "w"))
    print(f"db-export: {len(entries) - 1} doc(s) + manifest for {len(manifest['files'])} file(s) -> {len(plans)} write plan(s) in {a.out} (manifest last)")

def find_payloads(obj):
    if isinstance(obj, dict):
        if obj.get("kind") in ("pyqkit-doc", "pyqkit-manifest"):
            yield obj
            return
        for x in obj.values():
            yield from find_payloads(x)
    elif isinstance(obj, list):
        for x in obj:
            yield from find_payloads(x)

def db_collect(src):
    docs, manifest = {}, None
    for f in glob.glob(os.path.join(src, "**", "*.json"), recursive=True):
        try:
            obj = json.load(open(f))
        except (json.JSONDecodeError, UnicodeDecodeError):
            continue
        for p in find_payloads(obj):
            if p["kind"] == "pyqkit-manifest":
                if manifest is None or str(p.get("created_utc", "")) > str(manifest.get("created_utc", "")):
                    manifest = p
            else:
                docs[(p["path"], p["part"])] = p
    if not manifest:
        return {}, None, ["no manifest"]
    files, problems = {}, []
    for rel, meta in manifest["files"].items():
        parts = [docs.get((rel, i)) for i in range(len(meta["docs"]))]
        if any(p is None for p in parts):
            problems.append(f"{rel}: missing parts")
            continue
        data = "".join(p["text"] for p in parts).encode("utf-8")
        if sha256_bytes(data) != meta["sha256"]:
            problems.append(f"{rel}: sha mismatch")
            continue
        files[rel] = data
    return files, manifest, problems

def cmd_db_import(a):
    files, man, problems = db_collect(a.src)
    if not man:
        sys.exit("db-import: no manifest found - use the vault file route")
    if problems:
        print("db-import PROBLEMS: " + "; ".join(problems))
        sys.exit(1)
    install_files(files, a.vault)
    print(f"db-import: restored {len(files)}/{len(man['files'])} file(s) · batch {man.get('batch')} · all checksums verified")

# ------------------------------------------------------------------------------------------------
# status (plain language, three sections)
# ------------------------------------------------------------------------------------------------
def cmd_status(a):
    v = a.vault
    files = table(v, "l1/FILES.csv")
    inst = table(v, "l1/INSTANCES.csv")
    st = load_state(v)
    done = defaultdict(int)
    for f in files:
        if f["status"] == "DONE":
            done[f["year"]] += 1
    prog = open(rpath(v, "PROGRESS.md")).read() if os.path.exists(rpath(v, "PROGRESS.md")) else ""
    todo = [ln.strip()[6:] for ln in prog.splitlines() if ln.strip().startswith("- [ ]")]
    unread = [u for u in table(v, "l1/UNREADABLE.csv") if not u.get("recovered_from")]
    closed = ", ".join(map(str, st.get("years_closed", []))) or "none yet"
    print(f"Status - batch {st.get('batch', 0)} · {len(inst)} question rows saved · years closed: {closed}"
          + (f" · working on {st['year_in_progress']}" if st.get("year_in_progress") else "") + "\n")
    print("Done")
    print("\n".join(f"- {y}: {n} file(s) fully mined" for y, n in sorted(done.items(), reverse=True)) or "- Nothing mined yet.")
    print("\nNext (my work, in order)")
    print("\n".join(f"- {t}" for t in todo[:8]) or "- Waiting for the next year's files.")
    new = [f for f in files if f["status"] == "NEW"]
    if new:
        print(f"- {len(new)} received file(s) waiting to be mined.")
    print("\nOptional (yours)")
    print("\n".join(f"- A clear photo of {u['paper_id']} page {u['page']} would recover Q{u['q_no']} ({u['what_lost']})." for u in unread[:5])
          or "- Nothing needed right now.")

# ------------------------------------------------------------------------------------------------
# vault page (static; published once at Batch 0 with capabilities {db:{}, assets:{}})
# ------------------------------------------------------------------------------------------------
VAULT_PAGE = """<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">
<title>PHY-042 VAULT</title>
<style>
:root{--bg:#f6f4ef;--fg:#1d1d1b;--muted:#6b6862;--card:#ffffff;--line:#e2ded5;--accent:#2f5d50;
box-sizing:border-box;padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#141413;--fg:#ecebe6;--muted:#a19e96;--card:#1d1d1b;--line:#33312d;--accent:#8fc3b0}}
:root[data-theme="dark"]{--bg:#141413;--fg:#ecebe6;--muted:#a19e96;--card:#1d1d1b;--line:#33312d;--accent:#8fc3b0}
html{scroll-padding-top:env(safe-area-inset-top,0px);height:100%}
body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.55 Georgia,"Times New Roman",serif}
main{max-width:640px;margin:0 auto;padding:28px 20px}
h1{font:600 1.35rem/1.2 system-ui,-apple-system,"Segoe UI",sans-serif;margin:0 0 4px}
.k{color:var(--muted);font:500 .8rem system-ui,-apple-system,sans-serif;text-transform:uppercase;letter-spacing:.08em}
.card{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:16px 18px;margin:18px 0}
code{font:14px ui-monospace,Menlo,Consolas,monospace;color:var(--accent)}
</style></head><body><main>
<div class="k">Physics (042) · PYQ pattern mining</div>
<h1>PHY-042 VAULT</h1>
<div class="card">This page is the project's safe. Claude saves the mining ledger into it at every checkpoint and restores it
at the start of every chat.</div>
<div class="card"><b>Nothing to do here.</b> Don't delete this artifact. The newest <code>PHY-VAULT-…json</code> file in project
knowledge restores everything even without it.</div>
</main></body></html>
"""

def cmd_vault_page(a):
    sys.stdout.write(VAULT_PAGE)

# ------------------------------------------------------------------------------------------------
def main():
    try:
        import signal
        signal.signal(signal.SIGPIPE, signal.SIG_DFL)   # quiet when piped into head
    except (AttributeError, ValueError):
        pass
    ap = argparse.ArgumentParser(prog="pyqkit", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = ap.add_subparsers(dest="cmd", required=True)
    def p(name, fn, *args):
        s = sp.add_parser(name)
        for spec in args:
            s.add_argument(*spec[0], **spec[1])
        s.set_defaults(fn=fn)
    V = (["vault"], {})
    p("init", cmd_init, V)
    p("restore", cmd_restore, V, (["--db"], {"default": ""}), (["--sources"], {"nargs": "*"}))
    p("seed", cmd_seed, V, (["--syllabus"], {"default": ""}), (["--archetypes"], {"default": ""}))
    p("inventory", cmd_inventory, V, (["paths"], {"nargs": "+"}), (["--batch"], {"type": int}))
    p("pages", cmd_pages, (["doc"], {}), (["--vault"], {"default": ""}), (["--text"], {"type": int, "default": 0}))
    p("render", cmd_render, (["doc"], {}), (["pages"], {}), (["--vault"], {"default": ""}), (["--dpi"], {"type": int, "default": 200}),
      (["--crop"], {"action": "store_true"}), (["--pair"], {"action": "store_true"}),
      (["--region"], {"help": "fractions x0,y0,x1,y1"}), (["--out"], {"default": "/home/claude/render"}))
    p("merge", cmd_merge, V, (["part"], {}), (["--table"], {"required": True}))
    p("check-paper", cmd_check_paper, V, (["paper_id"], {}))
    p("similar", cmd_similar, V, (["--year"], {"type": int}), (["--cross"], {"action": "store_true"}))
    p("candidates", cmd_candidates, V, (["--instance"], {"default": ""}), (["--text"], {"default": ""}), (["--chapter"], {"default": ""}),
      (["--q-type"], {"dest": "q_type", "default": ""}), (["--any-chapter"], {"dest": "any_chapter", "action": "store_true"}),
      (["--k"], {"type": int, "default": 5}))
    p("validate", cmd_validate, V)
    p("derive", cmd_derive, V, (["--until"], {"type": int}))
    p("index", cmd_index, V)
    p("year-close", cmd_year_close, V, (["year"], {}), (["--xlsx-dir"], {"dest": "xlsx_dir", "default": "/mnt/user-data/outputs"}))
    p("export-xlsx", cmd_export_xlsx, V, (["year"], {}), (["--out"], {"default": "/mnt/user-data/outputs"}))
    p("plan-mock", cmd_plan_mock, V, (["--mode"], {"default": "likely", "choices": ["likely", "seed", "stress", "wildcard", "coverage", "chapter"]}),
      (["--seed"], {"type": int, "default": 1}), (["--chapter"], {"default": ""}), (["--out"], {"default": ""}),
      (["--tries"], {"type": int, "default": 20}))
    p("backtest", cmd_backtest, V, (["--year"], {"required": True}), (["--pred"], {"required": True}))
    p("pack", cmd_pack, V, (["--batch"], {"type": int}), (["--out"], {"default": "/mnt/user-data/outputs"}))
    p("unpack", cmd_unpack, (["file"], {}), (["--to"], {"required": True}))
    p("db-export", cmd_db_export, V, (["out"], {}))
    p("db-import", cmd_db_import, (["src"], {}), V)
    p("status", cmd_status, V)
    p("vault-page", cmd_vault_page)
    p("version", lambda a: print(f"pyqkit kit edition {KIT_EDITION} · schema edition {SCHEMA_EDITION}"))
    a = ap.parse_args()
    a.fn(a)

if __name__ == "__main__":
    main()
